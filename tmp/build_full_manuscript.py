from pathlib import Path
import csv
import json
from collections import defaultdict
import statistics
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
RUN=ROOT/'crave/runs/bdd2b8a990571bb1'
OUT=ROOT/'CRAVE_Full_Manuscript.docx'
ASSETS=ROOT/'tmp/full_manuscript_assets'
ASSETS.mkdir(exist_ok=True)
def read_csv(path):
    with open(path,encoding='utf-8') as f:return list(csv.DictReader(f))
summary=read_csv(RUN/'results/summary.csv')
diagnostics=read_csv(RUN/'results/value_diagnostic_summary.csv')
stats=read_csv(RUN/'results/statistics.csv')
raw=json.loads((RUN/'evaluation.json').read_text())['data']
audit=json.loads((RUN/'results/independent_audit.json').read_text())
compute=json.loads((RUN/'compute_accounting.json').read_text())
frozen=json.loads((RUN/'FROZEN_PILOT_CONFIG.json').read_text())
assert audit['passed'] and len(raw)==162
lookup={(r['dataset'],r['method']):r for r in summary}
methods=['CRAVE','confidence','fixed','llm_router','-CF','-RTG','-Interaction','-Budget','myopic']
names={'confidence':'Confidence rule','fixed':'Fixed schedule','llm_router':'LLM router','myopic':'Myopic alias'}
from expanded_manuscript_section import rewrite as expanded_rewrite,add_results as add_expanded_results
from scientific_manuscript_section import EXTRA_REFS,add_related,add_method,add_analyses
from journal_revision import rewrite as journal_rewrite,add_contributions,add_guarantees,add_influence,add_audit,add_interpretation
def n(row,key):return float(row[key])
def fmt(row,key,digits=1):return f'{n(row,key):.{digits}f}'
def interval(row,key,digits=1,scale=1):
    return f'{n(row,key)*scale:.{digits}f} [{n(row,key+"_ci_low")*scale:.{digits}f}, {n(row,key+"_ci_high")*scale:.{digits}f}]'

# Figures use the same frozen observations, with legible portrait-manuscript sizing.
if not all((ASSETS/name).exists() for name in ['quality_cost.png','value.png']):
    exec("""
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':9,'axes.spines.top':False,'axes.spines.right':False})
fig,axes=plt.subplots(2,1,figsize=(6.3,6.2))
colors=plt.get_cmap('tab10').colors
for ax,dataset in zip(axes,['HotpotQA','2WikiMultiHopQA']):
    for i,method in enumerate(methods):
        r=lookup[(dataset,method)]
        y=n(r,'exact_match')
        ax.errorbar(n(r,'generated_tokens'),y,yerr=[[y-n(r,'exact_match_ci_low')],[n(r,'exact_match_ci_high')-y]],
                    fmt='o',color=colors[i],capsize=2,label=names.get(method,method),ms=5)
    ax.set(xlabel='Mean generated output tokens including routing',ylabel='Exact match',title=f'{dataset}   Three task clusters',ylim=(-.06,1.06))
handles,labels=axes[0].get_legend_handles_labels()
fig.legend(handles,labels,loc='lower center',ncol=3,fontsize=8,frameon=False,bbox_to_anchor=(.5,-.005))
fig.tight_layout(rect=(0,.11,1,1))
fig.savefig(ASSETS/'quality_cost.png',dpi=250,bbox_inches='tight');plt.close(fig)
dc=read_csv(RUN/'results/counterfactual_diagnostics.csv')
d=[r for r in dc if r['method']=='CRAVE' and float(r['lambda'])==.1]
x=np.array([float(r['empirical_q']) for r in d]);y=np.array([float(r['predicted_q']) for r in d])
lo=min(x.min(),y.min())-.04;hi=max(x.max(),y.max())+.04
fig,ax=plt.subplots(figsize=(5.4,4.4))
ax.scatter(x,y,s=25,alpha=.7);ax.plot([lo,hi],[lo,hi],'k--',lw=1)
ax.set(xlabel='Finite K empirical utility under mu0',ylabel='Predicted utility',xlim=(lo,hi),ylim=(lo,hi),title='CRAVE value estimates at lambda 0.1')
fig.tight_layout();fig.savefig(ASSETS/'value.png',dpi=250);plt.close(fig)
""")

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.2677);sec.page_height=Inches(11.6929)
sec.top_margin=sec.bottom_margin=Inches(.78)
sec.left_margin=sec.right_margin=Inches(.85)
sec.footer_distance=Inches(.35)
for style in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
    s=doc.styles[style];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
    s._element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    s._element.rPr.rFonts.set(qn('w:ascii'),'Times New Roman');s._element.rPr.rFonts.set(qn('w:hAnsi'),'Times New Roman')
normal=doc.styles['Normal'];normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
normal.paragraph_format.widow_control=True
doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True
doc.styles['Title'].paragraph_format.space_after=Pt(7)
doc.styles['Subtitle'].font.size=Pt(12);doc.styles['Subtitle'].paragraph_format.space_after=Pt(12)
for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
    s=doc.styles[name];s.font.size=Pt(size);s.font.bold=True
    s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(5);s.paragraph_format.keep_with_next=True
doc.styles['Caption'].font.size=Pt(9.5);doc.styles['Caption'].font.italic=True
doc.styles['Caption'].paragraph_format.space_after=Pt(8)
for s in doc.styles:
    if s.type==1 and s.name in ['Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        borders=s._element.find('.//'+qn('w:pBdr'))
        if borders is not None:borders.getparent().remove(borders)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
footer.style='Normal';footer.paragraph_format.space_after=Pt(0)
doc.core_properties.title='CRAVE Counterfactual Value Learning for Budgeted Language Model Agents'
doc.core_properties.subject='Counterfactual resource allocation and an audited empirical pilot'
doc.core_properties.author=''
doc.core_properties.keywords='CRAVE; inference control; counterfactual branching; multi-hop question answering; pilot study'

def p(text,style=None):
    text=journal_rewrite(expanded_rewrite(text))
    v=doc.add_paragraph(text,style);v.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY if style is None else WD_ALIGN_PARAGRAPH.LEFT
    return v
def h(text,level=1):
    text=re.sub(r'^(\d+) (\d+) ',r'\1.\2 ',text)
    heading=doc.add_heading(text,level=level)
    if text in ['References','Appendix A Executable procedure']:
        heading.paragraph_format.page_break_before=True
def math_run(text,upright=False):
    run=OxmlElement('m:r')
    if upright:
        pr=OxmlElement('m:rPr');normal=OxmlElement('m:nor');pr.append(normal);run.append(pr)
    value=OxmlElement('m:t');value.text=text;run.append(value)
    return run
def math_part(part):
    if isinstance(part,str):return math_run(part)
    kind,base,index=part
    element=OxmlElement('m:sSub' if kind=='sub' else 'm:sSup')
    value=OxmlElement('m:e');value.append(math_run(base,upright=base=='arg max'));element.append(value)
    index_element=OxmlElement('m:sub' if kind=='sub' else 'm:sup')
    if isinstance(index,list):
        for child in index:index_element.append(math_part(child))
    else:index_element.append(math_run(index,upright=len(index)>1))
    element.append(index_element)
    return element
def eq(text,num):
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    v=doc.add_paragraph();v.alignment=WD_ALIGN_PARAGRAPH.CENTER
    v.paragraph_format.space_before=Pt(4);v.paragraph_format.space_after=Pt(7)
    S=lambda base,index:('sub',base,index)
    U=lambda base,index:('sup',base,index)
    formulas={
      1:[S('s','t'),' = (x, ',S('H','t'),', ',S('E','t'),', ',S('V','t'),', ',S('m','t'),', ',S('b','t'),')'],
      2:[S('A','f'),'(s) = {a ∈ A : ',S('r','j'),'(a) ≤ ',S('b','t,j'),', ∀ j ∈ J}'],
      3:[S('U','λ'),'(τ) = R − λ (',S('C','tok'),'/4096 + ',S('C','ret'),'/4 + ',S('C','ver'),'/2 + ',S('C','lat'),'/60)'],
      4:[U('Q','μ'),'(s, a; λ) = E[',S('U','λ'),'(τ) | s, a, μ]'],
      5:[U('Q̂','μ'),'(s, a; λ) = (1/K) ',S('∑','k=1,…,K'),' ',S('U','λ'),'(',S('τ','k'),'(s, a; μ))'],
      6:[U('I','μ'),'(s, a→b; λ) = ',U('J','μ'),'(a→b) − ',U('J','μ'),'(a) − ',U('J','μ'),'(b) + ',U('J','μ'),'(∅)'],
      7:['L = ',S('L','Q'),' + 0.25 ',S('L','cost'),' + 0.25 ',S('L','success'),' + 0.25 ',S('L','rank'),' + 0.25 ',S('L','interaction')],
      8:[S('π','θ'),'(s; λ) = ',S('arg max',['a ∈ ',S('A','f'),'(s)']),' ',S('Q','θ'),'(s, a; λ)'],
    }
    math=OxmlElement('m:oMath')
    for part in formulas[num]:math.append(math_part(part))
    v._p.append(math)
    v.add_run(f'    ({num})').font.size=Pt(10)
def caption(text):
    v=p(text,'Caption')
    v.paragraph_format.keep_with_next=text.startswith('Table')
    v.paragraph_format.keep_together=True
    return v
def table(headers,rows,widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');b.append(e)
    pr.append(b)
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,text in enumerate(headers):t.rows[0].cells[i].text=str(text)
    headerpr=t.rows[0]._tr.get_or_add_trPr();rep=OxmlElement('w:tblHeader');headerpr.append(rep)
    for values in rows:
        for i,text in enumerate(values):
            if i==0:cells=t.add_row().cells
            if str(text).startswith('Three untrained-on 2Wiki tasks only'):
                text='Twenty-three 2Wiki questions across two separate corpus protocols; full domain and cross-family model studies not run'
            cells[i].text=str(text)
    for ri,row in enumerate(t.rows):
        trpr=row._tr.get_or_add_trPr();unsplit=OxmlElement('w:cantSplit');trpr.append(unsplit)
        for i,cell in enumerate(row.cells):
            cell.width=Inches(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');margins.append(e)
            tcpr.append(margins)
            if ri==0:
                fill=OxmlElement('w:shd');fill.set(qn('w:fill'),'E7E7E7');tcpr.append(fill)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1);para.paragraph_format.space_before=Pt(1);para.paragraph_format.line_spacing=1.0
                para.paragraph_format.keep_with_next=ri<len(t.rows)-1
                para.alignment=WD_ALIGN_PARAGRAPH.LEFT if i==0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:run.font.size=Pt(9);run.font.bold=ri==0;run.font.color.rgb=RGBColor(0,0,0)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def figure(path,width,description):
    v=doc.add_paragraph();v.alignment=WD_ALIGN_PARAGRAPH.CENTER;v.paragraph_format.keep_with_next=True
    v.add_run().add_picture(str(path),width=Inches(width))
    description=re.sub(r'^Figure \d+',f'Figure {len(doc.inline_shapes)}',description)
    caption(description)

p('CRAVE Counterfactual Value Learning for Budgeted Language Model Agents','Title')
p('Matched interventions and resource allocation in multihop question answering','Subtitle')
p('Author names and affiliations to be supplied')
p('Method and empirical pilot study   October 2026')
h('Abstract')
p('Language model agents must decide whether to spend a limited inference budget on further reasoning, retrieval, verification, or termination. CRAVE learns the policy-conditioned long-horizon utility of these competing actions through matched counterfactual branching from identical intermediate states. A frozen text encoder and a compact multilayer perceptron predict action utility, eventual success, future resource costs, and auxiliary ordered interactions. A hard feasibility mask prevents token, retrieval, and verification overruns. We implement the framework using quantized Qwen3-4B and report an audited feasibility pilot rather than a completed benchmark study. Training used four HotpotQA and MuSiQue tasks, two collection rounds, 156 matched branch records, and K = 3 stochastic continuations. Evaluation comprised six held-out tasks, three seeds, and nine method labels, yielding 162 episodes. CRAVE achieved 33.3% exact match on three HotpotQA tasks and 66.7% on three 2WikiMultiHopQA tasks, matching a confidence rule while consuming 207.3 versus 62.4 and 262.6 versus 51.3 generated tokens, respectively. No observed efficiency advantage over this rule emerged. Independent checks verified state matching, split isolation, resource compliance, score recomputation, and artifact integrity. The pilot establishes an executable experimental method and documents negative findings; it does not establish general performance superiority, calibration, or transfer robustness. Full-data training, development tuning, broad budget sweeps, official literature baselines, and cross-model evaluation remain unexecuted.')
p('Keywords   Language models; inference budgets; counterfactual branching; adaptive retrieval; verification; resource allocation')

h('1 Introduction')
p('A language model agent can often improve an answer by generating another reasoning step, retrieving a missing fact, or checking a candidate. Each operation also consumes resources and can introduce new errors. The practical allocation problem is therefore sequential: the usefulness of an action depends on what the agent already knows, what it has tried, and which resources remain. Retrieval-augmented generation combines parametric knowledge with external evidence [1], while ReAct interleaves reasoning and environment actions [2]. These capabilities create a choice about when and how much to use each resource.')
p('Budget awareness and adaptive stopping already have substantial precedents. BudgetThinker conditions reasoning on remaining token capacity [3]. PUMA uses semantic redundancy and answer verification to identify candidate stopping points [4]. Confidence-based adaptive scaling couples confidence calibration with computation allocation [5]. TARG gates retrieval from a short retrieval-free draft [6]. These methods show that adaptive allocation is an established research area; adding a budget feature or a confidence threshold alone is insufficient to support a new contribution.')
p('The specific question studied here is whether a small controller can learn a comparable long-horizon value for several heterogeneous actions at the same intermediate state. A retrieval step may not immediately produce an answer but may make the next reasoning step useful. Verification may preserve a correct answer or replace it with an incorrect one. Additional reasoning may use evidence effectively or repeat an existing conclusion. A controller trained only on immediate rewards can miss these delayed consequences. A controller trained on observed trajectories alone may lack outcomes for actions that its collector rarely selects.')
p('CRAVE addresses these issues with controlled interventions within a fixed executor. For each recorded state, we restore the same state, force an eligible action, and execute a shared continuation policy. The resulting action values refer to that continuation policy and that executor. They are not policy-invariant causal effects or guaranteed values under the deployed controller. This distinction is central to the method and to the interpretation of its results.')
p('This manuscript presents the formulation, implementation, training objectives, and observed results of a bounded pilot. The completed experiment is deliberately small: four training tasks and six evaluation tasks. Its purpose is to establish that matched branching, learning, evaluation, and audit can run coherently on a local machine. Its main empirical finding is negative for efficiency: CRAVE matched the confidence baseline on answer accuracy and spent more generated tokens. The report retains this result rather than extrapolating favorable claims from the method design.')

add_contributions(p,h)
h('2 Related work')
h('2 1 Budget control and stopping',2)
p('BudgetThinker introduces control tokens and training procedures that make reasoning responsive to token limits [3]. PUMA identifies redundant reasoning and verifies proposed exits [4]. CAS uses calibrated confidence to allocate additional samples [5]. Their action families and supervision differ from CRAVE, but they constrain the novelty claim: reasoning length, answer readiness, and remaining budget are already useful allocation signals. The pilot confidence comparator is a fixed threshold rule over one candidate answer. It does not reproduce the calibration training and answer aggregation of CAS. Likewise, the implementation does not claim to reproduce the full PUMA protocol.')
h('2 2 Retrieval and heterogeneous actions',2)
p('TARG treats retrieval as a query-level decision derived from a no-context prefix [6]. DeepControl regulates information acquisition through retrieval continuation and granularity control [7]. Inference-Time Budget Control chooses search actions under token and tool limits using an operational value-of-information score and selective answer finalization [8]. These methods make adaptive information acquisition and heterogeneous budget accounting close prior work. CRAVE treats the retriever as a fixed module and asks how its estimated downstream return compares with reasoning, verification, and stopping at intermediate states.')
h('2 3 Counterfactual tool boundaries',2)
p('CoBRA learns the margin between internal and external tool-assisted experts using paired trajectories [9]. Its relevance is especially direct because the comparison is counterfactual and uses Qwen3-4B. CRAVE proposes repeated multi-action valuation within a trajectory, with common state restoration, explicit future resource costs, and auxiliary ordered interactions. This is a methodological distinction, not evidence that CRAVE outperforms CoBRA. None of these official systems was reproduced or evaluated in the completed pilot. Establishing comparative novelty and performance requires faithful implementations under a matched protocol.')

add_related(p,h)
h('3 Sequential resource allocation')
h('3 1 Observable state and actions',2)
p('The executor maintains the query x, bounded public summaries H, retrieved evidence E, a verification record V, monitoring features m, and remaining budget b. The observable state at step t is')
eq('sₜ = (x, Hₜ, Eₜ, Vₜ, mₜ, bₜ)',1)
p('The action set is A = {THINK, RETRIEVE, VERIFY, STOP}. THINK generates a candidate and a short public summary. RETRIEVE obtains additional passages. VERIFY checks the current candidate against available evidence using a separate instruction to the same model. STOP returns the current candidate without another answer-generation call. Verification has a distinct prompt, but its errors need not be independent of reasoning errors because the backbone is shared.')
p('Summaries and capped evidence are lossy representations of an evolving interaction. We therefore use an observable-state control formulation; we do not prove that the representation is a sufficient Markov statistic. The environment is finite because hard resource limits and a maximum of twelve executed actions bound each rollout. STOP is always eligible and can return an empty candidate; an empty output scores zero under the answer metric.')
h('3 2 Hard feasibility and actual consumption',2)
p('The feasibility mask reserves the maximum output cost of an action before execution. THINK requires at least 256 available generated tokens. VERIFY requires 192 tokens, one verification call, and a nonempty candidate. RETRIEVE requires one available retrieval call. STOP has no generated-token or tool reservation. The eligibility condition is')
eq('A_f(sₜ) = { a ∈ A : rⱼ(a) ≤ bₜ,ⱼ for each hard constrained resource j }',2)
p('Here J indexes the three hard constraints: output tokens, retrieval calls, and verification calls. The reservation rⱼ(a) is the maximum resource requirement used by the feasibility mask.')
p('The executor deducts actual generation counts after an eligible call completes. Reserving the upper bound is conservative: a short likely response can still be masked when fewer than 256 or 192 tokens remain. This restriction guarantees the implemented output-token cap without relying on a prediction of response length. It also changes the attainable frontier and must be shared across comparable methods. Input tokens and admitted context tokens are recorded separately and are not deducted from the pilot hard output budget.')
p('The executed budget was 512 generated tokens, one retrieval call, and one verification call per task. Latency was measured but not enforced as a hard deadline. A certified timeout executor would be required for a hard latency guarantee; the current environment explicitly rejects configurations that request one. Thus, the pilot does not implement every budget dimension of the broader proposed study.')
h('3 3 Utility and the policy conditioned estimand',2)
p('The terminal reward R is normalized answer exact match. Future cost includes all output tokens, retrieval calls, verification calls, and measured seconds from the intervention to termination. Reference scales were fixed before evaluation at 4096 tokens, four retrieval calls, two verification calls, and sixty seconds. For a continuation trajectory τ, the scalar utility is')
eq('Uλ(τ) = R − λ (Ctok/4096 + Cret/4 + Cver/2 + Clat/60)',3)
p('The action value is the expected utility after forcing action a from restored state s and then using continuation policy μ:')
eq('Qμ(s, a; λ) = E[ Uλ(τ) | restored state s, forced action a, continuation μ ]',4)
p('Returns are undiscounted. The cost is return-to-go, so an initially inexpensive action can receive a poor value if it induces expensive unsuccessful continuation. Conversely, an action with no immediate reward can receive value through a later correct answer. The stop value reflects the current candidate and the small measured stop overhead. No reward model or deployment-time access to gold answers is used.')

h('4 Matched counterfactual branching')
h('4 1 Restoration and randomized continuations',2)
p('A checkpoint serializes the complete observable state, resource counters, seed, step, and action log. Its content hash is verified on restoration. For each eligible intervention, the executor restores a fresh copy and checks equality with the pre-intervention checkpoint before assigning a randomized branch seed. Each replicate uses a seed derived from the recorded checkpoint hash, collector seed, and replicate index. Replicate seeds are common across alternative actions at that checkpoint, reducing unnecessary differences in randomization.')
p('The base backend receives an explicit, stateless prompt. Its internal token cache can accelerate repeated prefixes, but it does not add conversation information beyond that prompt. The inference state contains no gold answer, aliases, supporting-fact annotations, reward, or oracle action. Such fields remain in offline scorer records. Visited states are saved before their branches are generated, so resumption reuses the recorded intervention state instead of creating a new state with slightly different measured timing.')
p('For K stochastic continuations, the empirical target is')
eq('Q̂μ(s, a; λ) = (1/K) ∑ₖ Uλ(τₖ(s, a; μ))',5)
p('K was fixed at three for this pilot. A completed stability study would compare this choice with larger K on development tasks. Sample utility means, success means, and normalized cost means are targets; they are noisy estimates rather than exact values. Checkpoint identity is necessary for the matched interpretation, but it does not eliminate sampling variance or the dependence on μ.')
h('4 2 Label policy and visitation aggregation',2)
p('The bootstrap continuation μ0 assigns weights 0.35 to THINK, 0.35 to RETRIEVE, 0.20 to VERIFY, and 0.10 to STOP, renormalized over eligible actions. The STOP weight is reduced to 0.02 when no candidate exists. Policy randomness is reproducible from the state seed and step. The same μ0 was used to generate branch labels in both collection rounds.')
p('Round one visits the initial state and a state after THINK for each training task. A controller fitted to these labels then influences round-two visitation through a mixture that selects μ0 with probability 0.3 and the learned controller otherwise. A terminal decision can leave only an initial recorded state. The observed state counts were eight in round one and seven in round two. Changing visitation while keeping the label continuation fixed avoids silently changing the value estimand between rounds.')
p('The pilot produced 156 matched branch records and 111 records from naturally selected observational trajectories. Records from ordered pairs and empty interventions are included in the branch total. These counts are not 156 independent training tasks: the underlying questions number four, and several branches share a checkpoint and common replicate seeds. All learned controllers were trained only on training-split records.')
h('4 3 Ordered interactions',2)
p('To examine complementary actions, the branch machinery can force a feasible ordered pair, such as retrieval followed by reasoning, and compare its continuation utility with single-action and empty-intervention utilities:')
eq('Iμ(s, a→b; λ) = Jμ(a→b) − Jμ(a) − Jμ(b) + Jμ(∅)',6)
p('Jμ(z) denotes expected utility at the same state and utility weight after forcing the sequence z and then following μ. The empty sequence ∅ proceeds directly with μ without a forced action. Single-action J values coincide with the corresponding policy-conditioned Q values.')
p('Only replicate indices shared by all four intervention families enter the contrast. An ordered pair that becomes infeasible is excluded for the affected sequence. These contrasts depend on the continuation regime and on feasibility; they are not additive decompositions of a trajectory reward. The controller uses them as auxiliary shared-trunk supervision. They are not added again to the selected Q value, because the single-action return-to-go already includes downstream consequences. Interaction supervision was limited to one selected checkpoint per collection round, so broad complementarity claims are untested.')

add_guarantees(p,h)
h('5 Controller architecture and optimization')
h('5 1 Features and prediction heads',2)
p('The frozen all-MiniLM-L6-v2 encoder supplies 384-dimensional masked-mean-pooled and L2-normalized vectors [14]. Four vectors encode the query, the last three public summaries, accumulated evidence, and the candidate together with verification text. Encoder inputs are truncated at 256 tokens. Thirteen scalar features represent normalized remaining token, retrieval, and verification budgets; raw verbalized confidence; semantic convergence; evidence sufficiency; retrieval novelty; normalized step; utility weight λ; and a four-way previous-action indicator. The resulting input dimension is 1549.')
p('Confidence is the bounded value verbalized by the base model, not a calibrated probability. Convergence is summary embedding similarity. Novelty compares the embeddings of previous evidence and newly admitted evidence. Evidence sufficiency uses the bounded model support signal and is zero in the absence of evidence. These inexpensive proxies are operational definitions rather than validated measures of latent uncertainty or factual support.')
p('A two-hidden-layer ReLU multilayer perceptron has widths 512 and 256 and forty outputs. For each of four actions, six outputs encode Q, a success logit, and four future-cost logits. The success output is transformed by a sigmoid, and costs by softplus. Sixteen additional outputs represent ordered interaction contrasts. The base language model and encoder remain frozen; controller training uses NumPy analytic gradients and Adam on the CPU.')
h('5 2 Joint supervision',2)
p('The Q head uses Huber loss. The success head uses squared error against finite-K success proportions. Each resource head uses Huber loss against its normalized future cost. A pairwise logistic objective encourages ordering of actions with unequal empirical Q targets. Observed-target ties within 10⁻⁶ are excluded from the ranking objective. The overall pilot objective is')
eq('L = LQ + 0.25 Lcost + 0.25 Lsuccess + 0.25 Lrank + 0.25 Linteraction',7)
p('The pairwise logistic loss has the form log(1 + exp(d)) − y d, where d is the predicted action-value difference divided by temperature 0.1 and y indicates which empirical target is larger. The four auxiliary weights were fixed at 0.25. Adam used learning rate 0.001, batch size 64, one hundred epochs, gradient norm cap five, and seeds 17, 29, and 43. These defaults were not selected through development optimization.')
p('The training utility weights were 0, 0.05, 0.1, 0.25, 0.5, and 1.0. This gives the value predictor an explicit preference input, but it does not establish successful interpolation or generalization across preferences. Deployment evaluation used only λ = 0.1. A broad preference frontier cannot be inferred from this single operating point.')
h('5 3 Feasible action selection',2)
p('The deployed controller chooses the largest predicted Q among eligible actions:')
eq('πθ(s; λ) = arg max over a ∈ A_f(s) of Qθ(s, a; λ)',8)
p('The feasibility mask provides resource compliance independently of model prediction quality. A poor value estimate can still select a wasteful eligible action or stop too early. The successful budget audit therefore validates accounting and masking; it does not validate the learned allocation policy. Controller time is included in recorded latency, and LLM routing output is included in the output-token counter for the routing baseline.')

add_method(p,figure)
h('6 Experimental protocol')
h('6 1 Data and split boundaries',2)
p('HotpotQA supports multi-document factual and comparison questions with supporting facts [10]. MuSiQue composes connected single-hop questions into multihop tasks [11]. 2WikiMultiHopQA supplies reasoning-oriented questions and structured evidence [12]. The completed run used small first-record subsets, rather than random samples or full benchmark releases. They were chosen before model evaluation to make execution and audit feasible.')
caption('Table 1  Disjoint tasks and their executed roles')
table(['Dataset','Training','Reserved development','Original pilot','Expanded test'],[['HotpotQA','2','6','3','20'],['MuSiQue raw export','2','6','0','0'],['2Wiki raw export','0','0','3','20'],['Total','4','12','6','40']],[1.85,.8,1.35,1.15,1.4])
p('The expanded test was selected separately after the original pilot. Its forty questions are disjoint from the other roles in Table 1; Section 7.6 provides the frozen selection and evaluation protocol.')
p('HotpotQA records were obtained through the public dataset-server representation of the HotpotQA distractor data. MuSiQue training and development records, and 2Wiki development records, came from public raw exports linked by the MuSiQue authors. The subset files retain URLs and checksums. These raw exports are not asserted to be identical to complete official release files. The exact linkage between the dataset-server snapshot and the recorded repository revision also requires further verification for a definitive study.')
p('Task IDs, normalized questions, and available component-question identifiers were checked across partitions. The reserved development set was not used for threshold, K, hyperparameter, or loss-weight tuning in this run. The six evaluation tasks were excluded from learning and label aggregation. The 2Wiki tasks were absent from controller training, but this narrow transfer check does not establish out-of-domain robustness.')
h('6 2 Executor and corpus',2)
p('All compared policies use the official Qwen3-4B Q4_K_M GGUF checkpoint through llama.cpp release b11146 [13, 15]. Sampling uses temperature 0.7, top-p 0.8, and top-k 20. The executor requests bounded structured answers and public summaries in non-thinking mode. THINK permits at most 256 output tokens and VERIFY at most 192. This practical operation differs from unconstrained hidden-chain-of-thought continuation; results apply to this specified executor.')
p('The fixed corpus pools all 300 unique passages supplied by the selected training, development, and evaluation records, without reading supporting-fact flags or answers during corpus construction. Retrieval combines BM25 and frozen dense cosine rankings with reciprocal rank fusion constant sixty. Each RETRIEVE returns up to five documents. Newly admitted passages are capped at 280 tokens each and total added evidence at 1400 tokens. The query uses the question and the most recent available summary.')
p('Pooling supplied evaluation passages is a transductive benchmark-context setting. It is substantially easier and smaller than open retrieval over a full Wikipedia snapshot. The same pool is available to every policy, but corpus availability still limits external validity. The separate support-title evaluation reads gold annotations only after retrieval and never sends them to the agent.')
h('6 3 Comparators and ablations',2)
p('The fixed schedule requests RETRIEVE, THINK, VERIFY, THINK, and STOP, skipping masked actions. The confidence rule retrieves when no evidence exists, reasons while resources permit, and stops when a candidate has verbalized confidence at least 0.8. The LLM router uses the same backbone to select a permitted action in at most twenty-four output tokens; it reserves routing cost before executing that action. These are transparent in-house controls, not reproductions of named literature systems.')
caption('Table 2  Executed controller variants')
table(['Label','Change relative to CRAVE'],[['CRAVE','Matched long-horizon values with joint auxiliary objectives'],['−CF','Observed on-policy trajectory targets instead of alternate-action labels; interaction objective disabled'],['−RTG','Immediate reward and immediate resource cost targets; interaction objective disabled'],['−Interaction','Remove ordered interaction supervision'],['−Budget','Zero the remaining-budget scalar features; retain the external feasibility mask'],['Myopic alias','Same implemented learner as −RTG; duplicate comparator label']],[1.35,5.2])
p('The −CF comparison differs in action coverage and target quantity, not only in the presence of counterfactual labels. It therefore cannot isolate a clean causal-supervision effect. Myopic and −RTG have identical implementations and should not be counted as independent supporting experiments. The remaining proposed feature ablations, shuffled costs, independent-value baseline, and official TARG, PUMA, CAS, DeepControl, Budget Control, CoBRA, CATS, ATLAS, and LIMBO comparisons were not run.')
h('6 4 Metrics and statistical unit',2)
p('Answer exact match and token F1 use normalized strings and available aliases. For the raw exports, equivalence to each released benchmark evaluator and its complete alias mapping has not been established. Thus, the reported scores are pilot normalized answer metrics, not official leaderboard scores. Supporting-fact exact match or F1 is not reported as a task outcome.')
p('Results average the three seeds within each question before averaging questions. Bootstrap intervals resample task clusters with ten thousand draws. Paired tests operate on task-level differences and use exact sign permutations for these three-question dataset cells. Holm correction was applied across twenty-eight exploratory dataset, comparator, and metric comparisons. Seeds and episode rows are not treated as independent questions. With three task clusters per dataset, inferential resolution is extremely coarse.')
p('Raw output, input, and total tokens; retrieval and verification calls; logical latency; controller latency; and normalized composite cost are recorded. Cached generation results retain their original observed model latency for logical accounting, while physical request counts record work actually performed. Encoder caching, slot behavior, and action overhead prevent these logical times from being interpreted as a standardized cold end-to-end latency benchmark.')

h('6 5 AI assistance and analysis provenance',2)
p('OpenAI Codex with a GPT-6 family model assisted implementation, verification scripts, literature synthesis, manuscript drafting, and Matplotlib plotting code. Empirical plots are computed from saved numerical records; no image-generation model supplied result values or research images. The method schematic is a programmatic drawing. Exact Codex application and model build identifiers were not pinned, which limits reconstruction of the authoring process. The scientific workflow remains executable through the released analysis scripts and saved outputs. Human authors must independently review the methods, references, interpretation, and final text before approving submission.')
h('7 Results')
h('7 1 Answer quality and generated output cost',2)
p('Table 3 contains every executed method label at the single deployment setting. Each cell averages three seeds on three task clusters. CRAVE obtained exact match 0.3333 on HotpotQA and 0.6667 on 2Wiki. The confidence rule and fixed schedule obtained the same exact match. F1 equaled exact match for every method in this pilot. The LLM router returned no correct HotpotQA answers but matched the other retrieval policies on 2Wiki.')
caption('Table 3  Mean exact match and generated output tokens at a budget of 512 and utility weight 0.1')
rows=[]
for method in methods:
    a=lookup[('HotpotQA',method)];b=lookup[('2WikiMultiHopQA',method)]
    rows.append([names.get(method,method),f'{100*n(a,"exact_match"):.1f}',fmt(a,'generated_tokens'),f'{100*n(b,"exact_match"):.1f}',fmt(b,'generated_tokens')])
table(['Method','Hotpot EM percent','Hotpot output tokens','2Wiki EM percent','2Wiki output tokens'],rows,[1.65,1.18,1.24,1.18,1.3])
p('CRAVE used 207.3 generated tokens per HotpotQA task compared with 62.4 for the confidence rule, and 262.6 per 2Wiki task compared with 51.3. These ratios are approximately 3.32 and 5.11. The confidence rule therefore attained the same observed answer quality with lower output consumption in both dataset cells. No efficiency gain over that comparator is visible at the evaluated setting. This observation does not prove that the method cannot improve with adequate training or a different prespecified setting.')
p('Each nonzero accuracy cell has a task-bootstrap 95% interval spanning 0% to 100%. This reflects the small number of task clusters and the coarse binary outcomes. The paired accuracy difference between CRAVE and the confidence rule is exactly zero on the observed questions, giving a degenerate [0, 0] bootstrap interval for that difference. Neither interval addresses representativeness of first-record samples.')
h('7 2 Resource accounting',2)
caption('Table 4  Raw resource usage for the principal controls')
rows=[]
for dataset,short in [('HotpotQA','Hotpot'),('2WikiMultiHopQA','2Wiki')]:
    for method in ['CRAVE','confidence','fixed','llm_router']:
        r=lookup[(dataset,method)]
        rows.append([short,names.get(method,method),fmt(r,'input_tokens'),fmt(r,'total_tokens'),fmt(r,'verification_calls',2),fmt(r,'latency',2)])
table(['Dataset','Method','Input tokens','Total tokens','Verify calls','Logical seconds'],rows,[.72,1.43,1.10,1.15,.92,1.23])
p('Every principal control made one retrieval call per task on average. The fixed schedule made one verification call, whereas CRAVE, the confidence rule, and the LLM router made none. Input consumption was also higher for CRAVE than for the confidence rule: 4444.7 versus 1159.0 tokens on HotpotQA and 5906.0 versus 1059.0 on 2Wiki. Thus, the output-budget comparison does not conceal a compensating input-token advantage for CRAVE.')
p('CRAVE logical latency averaged 8.88 seconds on HotpotQA and 11.64 seconds on 2Wiki, compared with 3.82 and 2.89 seconds for the confidence rule. Its controller component averaged 0.198 and 0.223 seconds. The LLM router controller averaged 2.623 and 3.453 seconds. These values include the stated cache accounting convention and are descriptive measurements of this execution, not controlled deployment latency claims.')
p(f'The complete invocation executed {compute["physical_model_requests_this_invocation"]} physical model requests and generated {compute["physical_generated_tokens_this_invocation"]:,} physical output tokens. Recorded physical generation time was {compute["physical_generation_seconds_this_invocation"]:.1f} seconds and invocation time {compute["total_invocation_seconds"]:.1f} seconds, or about {compute["total_invocation_seconds"]/60:.1f} minutes. Logical evaluation output totaled {compute["logical_inference_tokens"]:,} tokens; logical branch return-to-go output totaled {compute["logical_offline_branch_generated_tokens"]:,}. The latter count overlaps cached computations and is not the physical training cost. Downloads, earlier component checks, and the backend smoke check are outside this invocation account. Monetary and energy costs were not measured.')
h('7 3 Paired exploratory comparisons',2)
caption('Table 5  CRAVE minus confidence rule on task-level paired observations')
rows=[]
for dataset in ['HotpotQA','2WikiMultiHopQA']:
    for metric in ['exact_match','normalized_cost']:
        r=next(r for r in stats if r['dataset']==dataset and r['comparison']=='CRAVE minus confidence' and r['metric']==metric)
        rows.append(['Hotpot' if dataset=='HotpotQA' else '2Wiki','EM' if metric=='exact_match' else 'Normalized cost',fmt(r,'estimate',4),f'[{n(r,"ci_low"):.4f}, {n(r,"ci_high"):.4f}]',fmt(r,'raw_p',2),fmt(r,'holm_p',2)])
table(['Dataset','Metric','Difference','95 percent interval','Raw p','Holm p'],rows,[.72,1.28,.92,1.75,.90,.98])
p('The normalized cost difference favors the confidence rule in both cells. Its bootstrap interval excludes zero on these observed first-record tasks, but the exact sign-permutation p value is 0.25. All twenty-eight Holm-adjusted p values equal 1.0. We therefore make no significance claim, no superiority claim, and no rejection of a broader research hypothesis. Resampling and correction cannot compensate for three task clusters or deterministic first-record selection.')
h('7 4 Held-out value diagnostics',2)
p('A separate analytical oracle branched from three held-out HotpotQA states after THINK, with K = 3 under μ0. Oracle labels were used only for diagnostics. Table 6 averages seed-level summaries at λ = 0.1. Regret is the finite-K empirical maximum utility minus the empirical utility of the action preferred by the learned Q head. Top-one agreement uses the same empirical maximum and its deterministic tie handling.')
caption('Table 6  Finite K held-out action-value diagnostics on three states')
rows=[]
for method in ['CRAVE','-CF','-RTG','-Interaction','-Budget']:
    selected=[r for r in diagnostics if r['method']==method and float(r['lambda'])==.1]
    keys=['mean_regret','top1_agreement','q_rmse','brier_to_finite_k_success_mean','ece_to_finite_k_success_mean']
    vals={k:statistics.mean(float(r[k]) for r in selected) for k in keys}
    rows.append([method,f'{vals[keys[0]]:.4f}',f'{100*vals[keys[1]]:.1f}',f'{vals[keys[2]]:.4f}',f'{vals[keys[3]]:.4f}',f'{vals[keys[4]]:.4f}'])
table(['Method','Regret','Top one percent','Q RMSE','Success MSE','Success ECE'],rows,[1.4,.85,1.1,1.0,1.05,1.15])
p('CRAVE mean regret was 0.0409, lower than −CF at 0.0652 but higher than −RTG at 0.0000 and −Interaction at 0.0226. Its top-one agreement was 44.4%. This pattern does not establish the intended return-to-go or interaction advantages. Indeed, the immediate-target model agrees with the empirical oracle on these states despite having zero answer accuracy in deployment. The diagnostic states have little measured opportunity for a successful continuation, so preferring inexpensive termination can score well in that limited diagnostic. Local oracle agreement and full task performance answer different questions.')
p('The success MSE is measured against finite-K success proportions, not individual Bernoulli outcomes. It omits the irreducible within-action success variance that an ordinary per-outcome Brier score would include. ECE likewise bins predictions against noisy finite-K proportions. These descriptive calibration proxies must not be reported as unbiased population calibration estimates. With K = 3, empirical utility maxima can be optimistic and action rankings unstable.')
h('7 5 Retrieval component check',2)
p('An independent question-only retrieval check evaluated top-five support-title coverage on the same pooled corpus. It included eleven HotpotQA tasks, eight MuSiQue tasks, and three 2Wiki tasks across their pilot roles. Gold support labels were consulted only by the offline evaluator. This component check did not change retrieval settings and did not supply targets to the deployed policy.')
ret=read_csv(ROOT/'crave/artifacts/retrieval_pilot/summary.csv')
caption('Table 7  Support-title recall in the separate retrieval component check')
rows=[]
for r in ret:rows.append([r['dataset'],r['n_tasks'],f'{100*n(r,"estimate"):.1f}',f'[{100*n(r,"ci_low"):.1f}, {100*n(r,"ci_high"):.1f}]',f'{100*n(r,"all_support_title_rate"):.1f}'])
table(['Dataset','Tasks','Recall percent','95 percent interval','All titles percent'],rows,[1.9,.65,1.15,1.5,1.35])
p('Mean support-title recall was 77.3% on HotpotQA, 62.5% on MuSiQue, and 83.3% on 2Wiki. All required support titles were retrieved for 54.5%, 25.0%, and 66.7% of tasks, respectively. A title match can conflate passages and does not verify that the relevant sentence was retained after token truncation. These values are a retrieval proxy, not official support-fact scores or evidence that the agent answered correctly.')

add_expanded_results(p,h,caption,table,figure)
add_analyses(p,h,figure)
add_influence(p,h,figure)
h('8 Audit and reproducibility')
p('Seventeen core tests passed, including atomic resource accounting, randomized budget safety, restore isolation, corrupt-checkpoint rejection, matched branch seeds, private-label exclusion, split checks, normalized metrics, analytic gradient finite differences, loss reduction, and statistical routines. Fixtures were used only in tests; research results came from the verified Qwen checkpoint and real dataset records.')
p('A separate audit recomputed configuration, source, stage, task, and corpus hashes; checked all 156 branch payloads and the 111 observational records; verified the fixed μ0 label regime; rejected duplicate records and private fields in inference states; recomputed held-out answers from gold records; reconciled output and total token counters; and checked token, retrieval, and verification bounds. It found no critical errors and zero budget violations in the 162 evaluation episodes.')
p('Artifact checks establish that the reported outputs correspond to the implemented pilot. They do not certify dataset release equivalence, official baseline fidelity, population-level effects, cross-model robustness, or cold latency. The immutable pilot configuration and source hash should be retained with any extension; a new experiment requires a new configuration and an explicit report of changed data or executor semantics.')

add_audit(p)
h('9 Discussion and limitations')
h('9 1 Interpreting the negative result',2)
p('The most direct completed comparison is against the confidence rule. It achieved the same answer accuracy with fewer input and output tokens in both dataset cells. More sophisticated action valuation did not produce better resource allocation at this pilot scale. Four training questions, fifteen visited states, noisy three-replicate targets, raw confidence features, and a fixed continuation regime provide plausible mechanisms for limited learning, but this experiment does not identify which mechanism caused the inefficiency. Explaining it requires prespecified development analysis and a larger independent evaluation.')
p('The framework may assign excessive value to continued reasoning when noisy training successes are not representative of held-out states. The original three-state diagnostic is consistent with optimistic predictions, but it does not establish a general overestimation mechanism. Similarly, a fixed λ and sparse budget support can leave a learned utility surface insufficiently constrained. The appropriate response is to test these explanations on development data rather than change evaluation settings after observing the held-out outcome.')
h('9 2 Statistical and data limitations',2)
p('Six questions cannot support a general benchmark conclusion. First-record selection is not a probability sample; three seeds do not increase the number of questions. Wide accuracy intervals and coarse permutation tests make population-level effects unidentifiable here. The raw exports and incomplete evaluator equivalence further restrict comparability. Pretraining contamination of the base model cannot be excluded from public factual QA data.')
p('The pooled corpus contains benchmark passages from evaluation questions. This avoids answer-label leakage into the policy but creates transductive evidence availability. A full Wikipedia retrieval study could yield different costs, support coverage, and action preferences. Multi-hop answer correctness alone also does not demonstrate connected reasoning, a concern that motivates the MuSiQue dataset design [11].')
h('9 3 Executor and estimand limitations',2)
p('Quantization, non-thinking mode, bounded public summaries, a single same-model verifier, capped evidence, and conservative token reservations define a particular action system. A controller evaluated with hidden-state features, longer internal reasoning, independent verification, or another backbone would face different transitions and costs. The pilot provides no invariance result across these choices.')
p('The learned values refer to μ0 continuation, whereas deployment uses the learned controller. This continuation mismatch is deliberate to preserve one label estimand across collection rounds, but it can limit policy improvement. Two rounds of changed visitation do not prove convergence. Ordered interaction targets use feasible common replicates and sparse checkpoint coverage; missing or ineligible pairs constrain what the auxiliary head can learn.')
p('Latency includes measured action and controller overhead together with the cache convention. It is useful for tracing this execution but is not a strict service deadline or a standardized hardware comparison. Neither deployment amortization of offline labels nor net financial savings has been established. The physical cost of generating labels remains relevant even when a learned policy eventually reduces online cost.')

add_interpretation(p,h)
h('10 Requirements for broader evaluation')
p('The broader design asks whether matched long-horizon values improve quality-cost frontiers, reduce regret compared with observational targets, capture delayed rewards, represent action complementarity, generalize to unseen budgets, and transfer across domains and backbones. Table 8 records the evidential status without treating successful execution as successful hypothesis confirmation.')
caption('Table 8  Research hypotheses and remaining evidence')
table(['Hypothesis','Status and required evidence'],[['H1 Quality and cost frontier','No efficiency gain over confidence at the tested point; full preference and budget frontier not run'],['H2 Counterfactual targets','Pilot diagnostics only; coverage-matched observational comparison and adequate task sample needed'],['H3 Delayed return','Immediate model performs poorly in task answers but well on three local oracle states; controlled delayed-benefit tasks needed'],['H4 Ordered interactions','Sparse auxiliary labels; no demonstrated gain; pair coverage and interaction analyses needed'],['H5 Unseen budgets','Not tested; one deployment budget executed'],['H6 Transfer','Three untrained-on 2Wiki tasks only; full domain and cross-family model studies not run']],[1.7,4.85])
p('A definitive follow-up should use complete, checksummed data releases and disjoint development-only selection of K, thresholds, loss weights, and hyperparameters. It should freeze a full retrieval snapshot, reproduce the closest literature baselines faithfully, match action availability and resource accounting, declare primary statistical comparisons, and evaluate adequate task counts across unseen budget combinations. Cross-model and larger-domain transfer should report calibration shift and failures, not just answer averages. These experiments are proposed requirements, not observations reported in this manuscript.')

h('11 Conclusion')
p('CRAVE implements matched, policy-conditioned long-horizon valuation of reasoning, retrieval, verification, and stopping under resource constraints. The pilot demonstrates reproducible collection, controller training, masked execution, real-data evaluation, and independent artifact auditing. It does not demonstrate an efficiency advantage: a simple confidence rule matched CRAVE accuracy with substantially fewer tokens on the six held-out questions. The contribution presently consists of the executable method and its transparent feasibility evidence. Claims about broad performance, causal-supervision benefits, interaction learning, budget generalization, and transfer remain open until the larger study is executed.')

h('Data and code availability')
p('Code, pinned configuration files, dataset subset provenance, raw branch and episode records, learned checkpoints, CSV summaries, statistical outputs, and validation artifacts are retained in the local CRAVE project. The principal result directory is crave/runs/bdd2b8a990571bb1. A public repository, archival identifier, and complete full-release dataset manifest have not been established. Public dataset and model sources are identified below and in Appendix B. Redistribution should follow the applicable upstream licenses; the 2Wiki raw-export license requires verification before a public artifact release.')
h('References')
refs=[
('Patrick Lewis et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS, 2020.','https://arxiv.org/abs/2005.11401'),
('Shunyu Yao et al. ReAct: Synergizing Reasoning and Acting in Language Models. ICLR, 2023.','https://arxiv.org/abs/2210.03629'),
('Hao Wen et al. BudgetThinker: Empowering Budget-aware LLM Reasoning with Control Tokens. arXiv:2508.17196v2, 2025.','https://arxiv.org/abs/2508.17196v2'),
('Dehai Min, Giovanni Vaccarino, Huiyi Chen, Yongliang Wu, Gal Yona, and Lu Cheng. Stop When Reasoning Converges: Semantic-Preserving Early Exit for Reasoning Models. arXiv:2605.17672, 2026.','https://arxiv.org/abs/2605.17672'),
('Xuqing Yang, Yi Yuan, Shanzhe Lei, and Xuhong Wang. Scaling with Confidence: Calibrating Confidence of LLMs for Adaptive Test Time Scaling. arXiv:2607.01612, 2026.','https://arxiv.org/abs/2607.01612'),
('Yufeng Wang, Lu Wei, and Haibin Ling. Retrieval as a Decision: Training-Free Adaptive Gating for Efficient RAG. arXiv:2511.09803v2, 2026.','https://arxiv.org/abs/2511.09803v2'),
('Siheng Xiong, Oguzhan Gungordu, James C. Kerce, and Faramarz Fekri. Adaptive Information Control for Search-Augmented LLM Reasoning. arXiv:2602.01672v2, 2026.','https://arxiv.org/abs/2602.01672v2'),
('Zhengru Fang et al. Inference-Time Budget Control for LLM Search Agents. arXiv:2605.05701, 2026.','https://arxiv.org/abs/2605.05701'),
('Wenhao Zou, Xianglong Liu, Wendong Bi, Hanjie Wang, Simin Zhao, and Gong Zhi. CoBRA: Learning Tool-Use Boundaries via Counterfactual Margins. arXiv:2609.00967, 2026.','https://arxiv.org/abs/2609.00967'),
('Zhilin Yang et al. HotpotQA: A Dataset for Diverse, Explainable Multi-hop Question Answering. EMNLP, pages 2369–2380, 2018. DOI 10.18653/v1/D18-1259.','https://aclanthology.org/D18-1259/'),
('Harsh Trivedi, Niranjan Balasubramanian, Tushar Khot, and Ashish Sabharwal. MuSiQue: Multihop Questions via Single-hop Question Composition. TACL, 10:539–554, 2022. DOI 10.1162/tacl_a_00475.','https://aclanthology.org/2022.tacl-1.31/'),
('Xanh Ho, Anh-Khoa Duong Nguyen, Saku Sugawara, and Akiko Aizawa. Constructing A Multi-hop QA Dataset for Comprehensive Evaluation of Reasoning Steps. COLING, pages 6609–6625, 2020. DOI 10.18653/v1/2020.coling-main.580.','https://aclanthology.org/2020.coling-main.580/'),
('Qwen Team. Qwen3-4B-GGUF official model repository. Revision bc640142c66e1fdd12af0bd68f40445458f3869b.','https://huggingface.co/Qwen/Qwen3-4B-GGUF'),
('Sentence Transformers. all-MiniLM-L6-v2 model card and implementation documentation. Revision 1110a243fdf4706b3f48f1d95db1a4f5529b4d41.','https://huggingface.co/sentence-transformers/all-MiniLM-L6-v2'),
('ggml-org. llama.cpp server documentation and Windows Vulkan release b11146.','https://github.com/ggml-org/llama.cpp/tree/b11146/tools/server'),
]
refs.extend(EXTRA_REFS)
for i,(text,url) in enumerate(refs,1):
    v=p(f'[{i}] {text} {url}')
    v.alignment=WD_ALIGN_PARAGRAPH.LEFT
    v.paragraph_format.space_after=Pt(3)
    v.paragraph_format.line_spacing=1.0
    for run in v.runs:run.font.size=Pt(9)

h('Appendix A Executable procedure')
p('Algorithm 1 specifies the labeling and deployment procedure used in the pilot. Gold labels are confined to the training scorer and held-out evaluator. The analytical oracle follows the labeling procedure on held-out states only after training.')
steps=[
'1  Construct and checksum task partitions and the common label-free passage corpus. Define action caps, budget rules, utility reference scales, and μ0 before evaluation.',
'2  For each training task, save an initial checkpoint and selected collector-visited checkpoints. Persist them before starting branch generation.',
'3  For each saved state and eligible action, restore and hash-check the full state, assign a common replicate seed, force the action, and continue with μ0 until STOP or the finite executor limit.',
'4  Score the terminal candidate offline and record actual future resource consumption. Repeat for K = 3. Generate feasible pair and empty-intervention records for selected interaction checkpoints.',
'5  Average branch targets for each state and action at each training λ. Create masks for missing actions and pair contrasts. Fit the shared controller on training records only.',
'6  Collect another visitation round with the μ0 and learned-controller mixture while retaining μ0 for every label continuation. Aggregate both rounds and fit the final seeded variants.',
'7  Evaluate fixed settings on held-out tasks. Select only feasible actions. Log all output, input, tool, and timing counters, final candidates, and action histories.',
'8  Recompute task metrics, pair observations by task, bootstrap task clusters, apply exploratory corrections, verify artifact checksums, and report negative as well as positive findings.'
]
for text in steps:p(text)
h('Appendix B Artifact identity and execution details')
p('The pilot configuration SHA256 is '+frozen['sha256']+'. The source SHA256 is '+frozen['config']['code_hash']+'. These identities describe the executed software and settings, not the later editorial manuscript builder.')
m=frozen['config']['model_manifest']
p('The Qwen file is '+m['model']['file']+' with SHA256 '+m['model']['sha256']+'. Its size was independently verified as 2,497,280,256 bytes before loading. The encoder ONNX SHA256 is '+m['encoder']['sha256']+'. The runtime archive SHA256 is '+m['runtime']['sha256']+'.')
p('The pooled passage corpus SHA256 is '+frozen['config']['corpus_hash']+'. The official source model and encoder revisions are listed in References 13 and 14. The raw export provenance points to https://raw.githubusercontent.com/StonyBrookNLP/musique/main/.all_data_information.json. Subset files record the author-linked public download URLs, subset checksums, and that full-file checksums were not obtained.')
p('Execution used Windows, Python 3.12.14, an RTX 3060 Laptop GPU with 6 GB dedicated memory, the Vulkan backend for Qwen generation, and CPU ONNX encoding. Dependency versions and the hardware setup record accompany the frozen run. The original pre-experimental Word manuscript is a separate preserved file. Reproduction commands and precise file formats are recorded in crave/RUNBOOK.md.')
h('Appendix C Interpreting the saved outputs')
p('per_instance.csv is the primary evaluation record. summary.csv contains seed-averaged task summaries and cluster-bootstrap intervals. statistics.csv contains paired differences, raw permutation probabilities, and Holm adjustments. counterfactual_diagnostics.csv retains action-level predictions, empirical utilities, and finite-K standard errors. value_diagnostic_summary.csv and success_reliability_bins.csv summarize diagnostic proxies. independent_audit.json and validation.json record distinct integrity checks. compute_accounting.json distinguishes physical inference during this invocation from logical token consumption.')
p('Checkpoint and raw trajectory files are essential for reviewing matched interventions. Learned model files are not evidence of generalization by themselves. The complete pilot record supports the numerical tables in this manuscript; unexecuted hypotheses, methods, domains, or budgets have no empirical result cells.')

appendix_identity=False
appendix_content=False
for paragraph in doc.paragraphs:
    if paragraph.text.startswith('Appendix A'):appendix_content=True
    if appendix_content and paragraph.style.name=='Normal':
        paragraph.paragraph_format.line_spacing=1.0
        paragraph.paragraph_format.space_after=Pt(3)
        for run in paragraph.runs:run.font.size=Pt(10)
    if paragraph.text.startswith('Appendix B'):appendix_identity=True
    if appendix_identity:paragraph.alignment=WD_ALIGN_PARAGRAPH.LEFT
doc.save(OUT)
plain='\n'.join(v.text for v in doc.paragraphs)
(ASSETS/'manuscript_text.txt').write_text(plain,encoding='utf-8')
print(json.dumps({'output':str(OUT),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),'words':len(plain.split()),'bytes':OUT.stat().st_size}))
