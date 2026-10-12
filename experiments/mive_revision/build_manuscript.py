"""Evidence-led major revision; preserve the source manuscript's native math.

Run with the bundled document runtime, only after the fresh audit verifies.
"""
import copy,csv,hashlib,json,pathlib,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=pathlib.Path(__file__).resolve().parents[2]
ART=ROOT/'experiments/mive_revision/artifacts';DEV=ART/'dev4b';FIG=ROOT/'experiments/mive_revision/figures'
TARGET=ROOT/'MIVE_Verified_Manuscript.docx'
SOURCE=pathlib.Path(r'C:/Users/ehsan/Downloads/CRAVE_Revised_Manuscript (2).docx')

def load(path):return json.loads(path.read_text(encoding='utf-8'))
def f(x):return 'undefined' if x is None else f'{x:.4f}'
def interval(row):
    if row['estimate'] is None:return 'undefined'
    return f"{f(row['estimate'])} [{f(row['ci95'][0])}, {f(row['ci95'][1])}]"

def main():
    assert load(DEV/'independent_verification.json')['passed']
    assert load(DEV/'statistics_verification.json')['passed']
    source_sha256=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    source=Document(SOURCE);doc=Document()
    section=doc.sections[0];section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=Inches(.8);section.bottom_margin=Inches(.8);section.left_margin=Inches(.9);section.right_margin=Inches(.9)
    for name in ['Normal','Body Text','Caption','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style.font.size=Pt(11)
        style.paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.1
    doc.styles['Body Text'].paragraph_format.line_spacing=1.1
    doc.styles['Title'].font.size=Pt(18);doc.styles['Title'].font.bold=True
    doc.styles['Subtitle'].font.size=Pt(11)
    for name,size in [('Heading 1',13),('Heading 2',11.5),('Heading 3',11)]:
        doc.styles[name].font.size=Pt(size);doc.styles[name].font.bold=True
        doc.styles[name].paragraph_format.space_before=Pt(12);doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE')
    field_run=OxmlElement('w:r');field_text=OxmlElement('w:t');field_text.text='1';field_run.append(field_text);field.append(field_run);footer._p.append(field)
    def p(text='',style=None):return doc.add_paragraph(text,style)
    def h(text,level=1):return doc.add_heading(text,level)
    def table(headers,rows,widths=None):
        t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        widths=widths or [6.7/len(headers)]*len(headers)
        for column,width in zip(t.columns,widths):column.width=Inches(width)
        for i,value in enumerate(headers):t.rows[0].cells[i].text=value
        for row in rows:
            cells=t.add_row().cells
            for cell,value in zip(cells,row):cell.text=str(value)
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
        t._tbl.tblPr.append(borders)
        for i,row in enumerate(t.rows):
            for j,cell in enumerate(row.cells):
                cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                pr=cell._tc.get_or_add_tcPr();margin=OxmlElement('w:tcMar')
                for edge in ['top','left','bottom','right']:
                    node=OxmlElement('w:'+edge);node.set(qn('w:w'),'75');node.set(qn('w:type'),'dxa');margin.append(node)
                pr.append(margin)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E4E9ED' if i==0 else ('F7F8FA' if i%2==0 else 'FFFFFF'));pr.append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(3);para.paragraph_format.space_before=Pt(3);para.paragraph_format.line_spacing=1
                    para.paragraph_format.keep_with_next=i==0
                    if i>0 and cell.text[:1] in '0123456789−-+':para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    for run in para.runs:run.font.size=Pt(9);run.font.bold=i==0
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        p()
        return t
    def cap(text):
        q=p(text,'Caption');q.paragraph_format.keep_together=True;q.paragraph_format.keep_with_next=text.startswith('Table');return q
    def figure(name,text,width=6.55):
        q=p();q.paragraph_format.keep_with_next=True;q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        picture=q.add_run().add_picture(str(FIG/(name+'.png')),width=Inches(width));picture._inline.docPr.set('descr',text);cap(text)
    def native(parts,number):
        q=p();q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math=OxmlElement('m:oMath')
        def run(text):
            r=OxmlElement('m:r');v=OxmlElement('m:t');v.text=text;r.append(v);return r
        for part in parts:
            if isinstance(part,str):math.append(run(part))
            else:
                kind,base,index=part;node=OxmlElement('m:sSub' if kind=='sub' else 'm:sSup')
                e=OxmlElement('m:e');e.append(run(base));node.append(e)
                e=OxmlElement('m:sub' if kind=='sub' else 'm:sup');e.append(run(index));node.append(e);math.append(node)
        q._p.append(math);q.add_run('    ('+str(number)+')')
    S=lambda a,b:('sub',a,b);U=lambda a,b:('sup',a,b)
    def retain(index,text=None):
        original=source.paragraphs[index]
        node=copy.deepcopy(original._p)
        doc.element.body.insert(len(doc.element.body)-1,node)
        result=doc.paragraphs[-1]
        if text is not None:result.text=text
        if original.style.name.startswith('Heading'):
            result.style=doc.styles[original.style.name]
        if text is not None:result.style=doc.styles['Heading '+str(text.split()[0].count('.')+1)]
        # Native math and source run content stay intact; clear incompatible colors.
        for run in result.runs:
            run.font.color.rgb=RGBColor(0,0,0)
            if not result.style.name.startswith('Heading'):run.font.size=Pt(11)
        return result

    summary=load(DEV/'summary.json');secondary=load(DEV/'utility_summary.json');verification=load(DEV/'independent_verification.json')
    cfg=load(DEV/'freeze.json');counts=summary['counts'];k=summary['K_selection']['reported_K']
    archive=load(ART/'archival_reanalysis.json')
    completion=load(DEV/'completion.json')
    gate=summary['variance_gate'];ks=summary['K_selection']
    p('Learned action allocation consumes more tokens without reliable accuracy gains in a budgeted language model agent','Title')
    p('Matched intervention values and continuation sensitivity in MIVE and CRAVE','Subtitle')
    p('Ehsan Derakhshi Ghazijahani')
    p('Politecnico di Torino, Turin, Italy\nCorrespondence: ehsan.derakh@gmail.com\nORCID: 0009-0009-4394-5948')
    p('12 October 2026\nArticle type: methodological analysis and diagnostic evaluation','Body Text')
    h('Abstract')
    p('Budgeted language model agents must decide whether to reason, retrieve, verify, or stop. Matched intervention value estimation (MIVE) compares these actions from an identical restored executor state under a declared continuation policy. We distinguish feasibility guarantees, finite-sample uncertainty, and continuation error. A frozen-controller evaluation of 40 questions from bounded HotpotQA and 2WikiMultiHopQA exports comprises 480 episodes across three seeds. Legacy CRAVE achieved normalized exact match of 0.5500 and 0.4333, versus 0.5000 and 0.4667 for a confidence rule, while using 229.5167 and 249.6833 output tokens versus 51.2000 and 60.0500. Accuracy differences were not significant after the original Holm adjustment; token costs were higher in both datasets. [E1–E3] '
      f'A fresh development audit on {counts["questions"]} questions from HotpotQA, 2WikiMultiHopQA, and MuSiQue compares the original labeling continuation with the deployed legacy controller using disjoint A/B replicate sets. It retained {counts["branches"]} branches from {counts["checkpoints"]} checkpoints. Signed THINK bias was {interval(summary["signed_THINK_bias"])} and the mean absolute discrepancy was {interval(summary["absolute_discrepancy"])} on the terminal-EM scale. [E4–E7] '
      +('The prespecified K stability criterion failed, so these finite-K discrepancy estimates remain exploratory. ' if ks['selected'] is None else f'The development stability rule selected K = {ks["selected"]}. ')
      +'We add a finite-sample ranking bound and show why raw EM alone cannot produce a strict THINK-versus-STOP reversal at a fixed checkpoint. Resource-priced values are therefore analyzed separately. The evidence establishes an allocation failure of this legacy controller and an executable continuation audit; it does not establish that refreshed MIVE training improves held-out efficiency.')
    p('Keywords: language model agents; action value learning; continuation policy; matched interventions; resource budgets; retrieval gating.','Body Text')
    h('1 Introduction')
    p('An agent can obey a token cap and still allocate its computation poorly. Reasoning, retrieval, verification, and stopping alter different parts of the execution state and consume different resources. A useful controller needs values that compare their consequences under the policy that will actually complete the trajectory. A terminally correct answer is insufficient to determine whether another reasoning call was worth its cost.')
    p('Restoring an intermediate state permits an executable comparison: force each eligible action and run a declared continuation. The resulting action value is conditioned on the model, prompt, evidence index, budget, state, and continuation. State matching removes variation in the starting point. It does not remove Monte Carlo noise or make a value learned under one continuation valid under another.')
    p('This distinction matters in the retained CRAVE evaluation. Relative to a confidence controller, the learned legacy controller increased generated tokens across both evaluated datasets without a statistically reliable accuracy gain. [E1–E3] That finding motivates a continuation audit, but it does not identify continuation mismatch as the cause. Training scarcity, value approximation, and stopping behavior remain competing explanations.')
    p('The contributions are a precise policy-conditioned estimand, conditional hard-budget and decision-loss results, a new finite-sample ranking bound, and a development audit with independently sampled continuations and a separate randomness-coupling comparison. The revised empirical claim concerns the observed legacy allocation failure. Optional continuation refresh is a proposed remedy whose held-out benefit remains to be established.')
    h('2 Related Work')
    p('Adaptive Computation Time [16] and PonderNet [17] learn when internal computation should stop. BudgetThinker [3] controls reasoning length through budget-aware training. Recent work uses semantic convergence [4], calibrated confidence [5], information control [7], or explicit search budgets [8]. These approaches motivate computation allocation, but their signals and supervision differ from action values estimated after restoring a complete agent state.')
    p('Retrieval-augmented generation [1] and ReAct [2] establish evidence use and interleaved reasoning as useful agent operations. Self-consistency [20] spends additional samples to aggregate answers, whereas action allocation asks which operation should receive the next resource unit. Long-context position effects [21], confidence calibration [22], and semantic uncertainty [23] further motivate checking the information available to a stopping rule. This study records a generator-provided confidence scalar; it does not establish its calibration.')
    p('Retrieval gating also includes learned decisions. Adaptive-RAG [27] selects among retrieval strategies using a complexity classifier. Self-RAG [19] learns retrieval and critique through reflection tokens, while FLARE [28] uses low-confidence future text to trigger retrieval. Retrieval as a Decision [6] studies training-free gating. CoBRA [9] learns tool-use boundaries from counterfactual reward margins. These systems are substantive comparators; replacing their trained components with an arbitrary confidence threshold would not reproduce them.')
    p('DAgger [18] addresses the distribution of visited states, and fitted Q iteration [29] relates value fitting to sequential decision making. Changing visitation alone does not refresh the continuation behind an old return label. MIVE makes that policy identity explicit. Table 1 compares estimands and accounting requirements, rather than claiming novelty from the use of branching alone.')
    p('The resource-constrained formulation follows constrained Markov decision processes [30]. Retriever choices such as Contriever [31] and ColBERTv2 [32] alter the evidence transition kernel, making the corpus and retriever part of a value target’s identity. They are related retrieval architectures, not evaluated components of this experiment.')
    cap('Table 1. Closest related methods and the distinctions relevant to a matched-budget comparison. Conceptual comparison from primary sources [3–9, 19, 27, 28]; no reproduced performance is implied.')
    table(['Work','Supervision or estimand','Continuation handling','Cost accounting needed'],[
      ['Adaptive-RAG [27]','Query complexity and strategy labels','Strategy determines later retrieval','Classifier, generation, every retrieval'],
      ['Self-RAG [19]','Reflection and critique token training','Native trained generator','Reflection tokens, retrieved text, generation'],
      ['FLARE [28]','Future-token confidence trigger','Draft and regenerate after retrieval','Discarded drafts and all retrieval stages'],
      ['BudgetThinker [3]','Budget-conditioned reasoning reward','Budget-aware trained backbone','Training and all inference tokens'],
      ['CoBRA [9]','Counterfactual internal versus tool reward margin','Expert/reference-split trajectories','Paired labels, training, tool calls'],
      ['MIVE in this work','Action return at a restored state','Explicit labeling and deployment policies','Offline branches plus deployed resource ledger']],[1.1,1.9,1.8,1.9])

    h('3 Problem and Method')
    headings={21:'3.1 State and action semantics',27:'3.2 Reservations and feasible actions',35:'3.3 Expected resource targets',44:'3.4 Policy conditioned action values',51:'3.5 Matched intervention value estimation',52:'3.5.1 Checkpoint restoration',58:'3.5.2 Collection and learning',63:'3.6 Common randomness and contrast variance'}
    for index in range(21,68):
        if index in [61,62]:continue
        retain(index,headings.get(index))
    p('For the executable legacy controller, four frozen text embeddings and budget/confidence scalars feed a two-hidden-layer network. Its scalar action head, not the proposed composite redesign, selects the largest feasible value. The generator and encoder remain frozen. All available controller identities and training records are retained; this revision does not relabel that legacy checkpoint as a newly trained MIVE policy. [E1, E4]')
    figure('workflow','Figure 1. Independent continuation measurement and separate sampler-seed coupling audit. The protocol applies to the fresh HotpotQA, 2WikiMultiHopQA, and MuSiQue development study with Qwen3-4B Q4_K_M; one prefix seed (17), independent A/B replicate sets, and a separate matched-seed matrix. Checkpoint and branch counts are reported in Section 6.2. [E4–E7]')

    h('4 Theory')
    headings={69:'4.1 Signed and absolute discrepancy',76:'4.2 Continuation shift and constructive reversal',84:'4.3 Approximation and decision error'}
    for index in range(69,89):
        if index==83:continue
        if index==81:
            p('Constructive example. Consider an incorrect stored candidate with deterministic EM zero and four available nonterminal action slots. Each THINK adds no information and has priced cost 0.1. A recovery operation produces a correct answer at priced cost 0.1 and terminates. After a forced THINK, labeling continuation μ performs recovery, giving return 1 − 0.1 − 0.1 = 0.8. Deployment continuation π repeats THINK in the three remaining slots and returns the unchanged candidate, giving return 0 − 4 × 0.1 = −0.4. STOP returns zero. Thus THINK outranks STOP under μ and ranks below it under π. The example is analytical; these values are not benchmark observations.')
            continue
        retain(index,headings.get(index))
    p('The constructive reversal concerns resource-priced return. For unpriced terminal EM, STOP at a fixed saved state has deterministic value zero or one. THINK has expected EM in the unit interval, so its difference from STOP cannot be positive under one continuation and negative under the other. A strict reversal count of zero for raw EM therefore cannot disprove continuation mismatch. The secondary priced analysis below tests the constructive example on the appropriate scale.')
    h('4.4 Finite sample ranking with a deterministic STOP action',2)
    p('Theorem 3 (sampling and continuation margin). Fix N checkpoints and a positive integer K before drawing audit replicates. Let R, ε, and m be positive, and let α lie strictly between zero and one. At each checkpoint, let independent THINK returns under a fixed labeling continuation be identically distributed in an interval of width R. STOP returns its current candidate without additional charge and is deterministic. With probability at least one minus α, every THINK mean differs from its labeling-continuation expectation by at most ε whenever')
    native(['K ≥ (',U('R','2'),' / (2',U('ε','2'),')) log(2N / α).'],15)
    p('If the true THINK-versus-STOP gap has absolute value at least m, the probability of any incorrect sign among the N fixed checkpoints is at most')
    native(['2N exp(−2K',U('m','2'),' / ',U('R','2'),').'],16)
    p('If THINK continuation discrepancy is additionally bounded in absolute value by δ at a checkpoint, an observed labeling gap greater than δ plus ε certifies a positive deployment gap on that event. The analogous negative-margin condition certifies stopping. This is a sufficient condition; failure to meet it is inconclusive.')
    p('Proof. Hoeffding’s inequality [35] bounds the two-sided deviation of each THINK sample mean by twice exp(−2Kε²/R²). A union bound over the fixed checkpoints gives (15). An incorrect sign with a true gap of magnitude at least m requires a deviation of at least m; applying the same bound gives (16). STOP is exact, so no STOP sampling term is needed. Finally, subtract at most δ of continuation shift and ε of sampling error from the observed gap. This establishes the margin statement. If STOP is stochastic or incurs a random finalization cost, its uncertainty must also be included; the stated theorem then does not apply unchanged.')
    p('The hard-resource ledger and deterministic STOP semantics are checked in the fresh branch reconstruction. The independent sets use disjoint sampler seeds, which operationalizes replication but does not prove independence of a pseudorandom backend. The uniform continuation bound required for a population policy-improvement certificate is not empirically established. For the secondary objective, latency is excluded and the return-range bound follows from the generated-token, retrieval, and verification caps. [E4, E7, E8]')
    p('The displayed bound fixes both checkpoints and K before the audit draws. Adaptive checkpoint fallback and empirical K selection are not treated as a certificate under this fixed-design statement. For a finite predeclared set of candidate K values, a further union bound replaces N by N times the number of candidates, provided the full candidate checkpoint set is fixed independently of those audit draws. The reported bootstrap intervals do not assert this simultaneous guarantee.')

    h('5 Experimental Setup')
    p('The completed experiments are a bounded, quantized Qwen3-4B diagnostic study [13–15]. The retained evaluation covers HotpotQA [10] and 2WikiMultiHopQA [12] at one resource cap; the fresh development measurement adds MuSiQue [11] and independent continuation replicates. These are distinct evidence blocks with distinct question and corpus manifests. Results are not pooled across them. [E1, E4]')
    h('5.1 Retained frozen controller evaluation',2)
    p('The legacy controller was fitted using four training questions and 156 matched branches, with a fixed exploratory continuation in both collection rounds. Its post-pilot evaluation froze controller identities, the confidence threshold (0.8), generation settings, and resource semantics before inference on 20 fresh questions per dataset. Three seeds (17, 29, 43), four methods, and a cap of 512 generated tokens, one retrieval, and one verification produced 480 episodes. The comparison used a common pool of 687 supplied passages, not full Wikipedia. These numerical records were independently reconstructed in this revision. [E1–E3]')
    p('The legacy methods are CRAVE, confidence, a fixed retrieve–think–verify–think–stop schedule, and an LLM router. Router output is included in the generated-token ledger. THINK and VERIFY reserve 256 and 192 output tokens respectively. Actual charges release unused capacity. Input tokens, admitted context, controller time, and total recorded latency are reported separately; input tokens and elapsed time are not enforced as hard caps. Cached legacy responses retain their original measured latency, precluding a cold-runtime speed comparison. [E1, E2]')
    p('The legacy scalar training target additionally prices recorded duration/60 together with generated tokens/4096, retrievals/4, and verifications/2 at weight 0.1. This empirical target has no hard latency bound. The fresh primary EM audit and bounded secondary objective are declared separately; the bounded-return theory is not a certificate for uncontrolled legacy latency. [E1, E4, E8]')
    h('5.2 Fresh development measurement',2)
    p(f'The fresh audit fixes {counts["questions"]} development questions, three per dataset, chosen by a deterministic hash order with selection seed 20261011 after excluding every retained pilot and expanded evaluation ID. The complete task and supplied-context corpus manifests were frozen before successful collection. Prefix generation uses seed 17. The exact Qwen3-4B model and seed-17 legacy controller are identified by SHA256. The runtime uses the local RTX 3060 Laptop GPU and four CPU inference threads. Collection began with one 4096-token slot and resumed with four 4096-token slots after a recorded runtime amendment. All serial records were retained, and the tasks, budgets, seeds, K candidates, and gate thresholds were unchanged. Generation-response caching is disabled. Setup failures and configuration supersessions are retained in the experiment register. [E4, E9]')
    p('The labeling continuation is the original feasible mixture with weights 0.35 THINK, 0.35 RETRIEVE, 0.20 VERIFY, and 0.10 STOP; STOP receives weight 0.02 when no candidate exists. The deployment continuation is the frozen legacy CRAVE controller at price 0.1. Checkpoints initially follow a THINK call without retrieval, leaving retrieval available for later recovery. A frozen fallback adds checkpoints following RETRIEVE then THINK only if more than 90% of initial checkpoints have identical K=5 action-mean EM. This is development checkpoint selection, not held-out tuning. [E4]')
    p('The checkpoint digest is verified before assigning an exogenous branch-seed key. This key selects replicate randomness and is excluded from controller features; question, candidate, history, evidence, counters, and remaining resources are restored identically. The seed override does not represent a new observed decision state. [E4, E7]')
    p('For every feasible action, each continuation receives separate A and B sets of five branches. Seed namespaces distinguish action, continuation, set, checkpoint, and replicate. The stability rule checks K=3 prefixes and then K=5: both continuations must have at least 90% A/B argmax agreement and at least 90% of top-two margins with Welch 95% half-width at most 0.10 in both sets. A failed rule leaves the K=5 analysis exploratory. Zero empirical variance can produce a zero plug-in Welch width and does not certify precision for rare outcomes. [E4, E5]')
    p('These Welch widths describe sample-selected top-two margins; they are not simultaneous post-selection confidence guarantees. Passing this operational rule selects a development replication count, not a population ranking certificate. [E4–E6]')
    p('The primary estimand is expected terminal normalized answer EM. The secondary value subtracts 0.1 times future output tokens/4096, retrieval calls/4, and verification calls/2. This sensitivity analysis was declared during development collection, before its first complete branch matrix; it is not represented as part of the initial freeze. Latency is recorded but excluded from this bounded secondary objective. All raw reward and cost records remain unchanged. [E8]')
    p('The primary EM stability rule selects the displayed K for both objectives. It does not certify stability of the secondary utility rankings. Changing server concurrency may alter floating-point generation behavior despite fixed sampler seeds; the execution version is retained per branch and exact cross-runtime replay is not assumed. [E4, E8, E9]')
    p('Secondary returns are calculated in exact integer units of 1/40960 before averaging, avoiding arithmetic-order artifacts in tied ranks. This implements the unchanged declared prices and does not modify raw answers or resource counts. The implementation clarification is retained beside E8.')
    p('Independent reconstruction detected that the initial parallel analysis paired matched branches by journal order. All paired contrasts were corrected to align replicate IDs; raw answers, seeds, resource counts, and frozen collection code were preserved. The initial derived-analysis hashes and correction record are retained in pairing_correction.json beside E5.')
    p('Kendall tau-b accounts for tied ranks; all-tied cases are undefined and their denominator is reported. Argmax uses the fixed action order THINK, RETRIEVE, VERIFY, STOP for ties. Strict rank reversal excludes ties. Behavioral disagreement is computed separately as one minus the labeling mixture probability of the legacy controller’s chosen action at a checkpoint. It measures policy behavior rather than disagreement of estimated value maximizers. [E5–E8]')
    p('Signed THINK bias is the labeling value minus the deployment value; a positive value indicates optimistic labeling. The experimental mean absolute discrepancy averages absolute differences over feasible actions. The theoretical discrepancy in (11) is a maximum over actions. Neither a finite-K average nor a sampled maximum supplies the uniform bound required by the policy-loss result. [E5, E8]')
    p('For each action, disjoint A/B contrasts estimate squared shift as (μA−πA)(μB−πB). Sampling noise is estimated from the within-continuation A/B differences. The shift estimator can be negative in a finite sample and is retained without truncation. These components decompose the expected squared discrepancy; they need not add exactly in one finite realization. Reported intervals use 10,000 question-cluster bootstrap resamples, weighting checkpoints equally within each question. They condition on the observed finite-K branch estimates and do not remove branch noise. [E5, E6, E8]')
    p('The separate coupling audit uses shared native sampler seeds by generation-call position across forced-action branches. It compares THINK-versus-alternative contrast variances with independent seeds at equal K per action and continuation. This implements a declared seed coupling; it does not assert identical random draws across divergent token sequences or hardware. Equal branch counts need not imply equal numbers of physical generation calls, which are reported separately. [E4–E8]')
    h('5.3 Scoring and inferential scope',2)
    p('Answer EM and token-overlap F1 use the retained normalization and supplied aliases. Official-release identity and complete entity-alias equivalence have not been established for the remapped bounded exports; scores are therefore labeled normalized answer metrics rather than official leaderboard scores. Quality intervals cluster questions after averaging seeds [24]. The original held-out permutation tests [26] and their Holm family [25] are retained as historical analyses, not repurposed as noninferiority tests. [E1–E3]')

    h('6 Results')
    h('6.1 Legacy allocation uses more generation',2)
    cap('Table 2. Retained evaluation quality and generation. HotpotQA and 2WikiMultiHopQA; Qwen3-4B Q4_K_M; expanded evaluation split; 20 questions per dataset and three seeds. Means and question-bootstrap 95% intervals are recomputed in E3; complete raw outcomes and original testing artifacts are in E1–E2.')
    rows=[]
    for r in archive:
        m=r['metrics'];em=m['exact_match'];tokens=m['generated_tokens']
        rows.append([r['dataset'].replace('MultiHopQA',''),r['method'],f(em['mean']),f(m['f1']['mean']),f"{f(tokens['mean'])} [{f(tokens['ci95'][0])}, {f(tokens['ci95'][1])}]"])
    table(['Dataset','Method','EM','F1','Output tokens [95% CI]'],rows,[1.0,1.1,.65,.65,3.3])
    p('The legacy CRAVE-minus-confidence EM differences were +0.0500 on HotpotQA and −0.0333 on 2WikiMultiHopQA. The original paired question-bootstrap intervals were [0.0000, 0.1500] and [−0.1000, 0.0000], and both original Holm-adjusted permutation p values were 1.0000. Additional generated tokens were 178.3167 [149.1154, 203.9171] and 189.6333 [165.9163, 211.3508], with adjusted p = 0.0018 in both datasets. These are tests of the historical contrasts, not evidence of equivalence or noninferiority. Each accuracy difference depended on a single question. [E1–E3]')
    figure('archival_quality_cost','Figure 2. Quality and generated-token operating points for the four legacy policies on HotpotQA and 2WikiMultiHopQA. Qwen3-4B Q4_K_M, expanded evaluation split, 20 questions per dataset, three seeds (17, 29, 43), one cap of 512 generated tokens. Error bars are marginal 95% question-bootstrap intervals from E3. These points do not constitute a multi-cap Pareto frontier.')
    cap('Table 3. Additional charged resources in the same retained evaluation as Table 2. Input tokens and latency are recorded resources, not hard caps; latency includes retained cache measurements. Values are means from E3 and raw records in E1.')
    table(['Dataset','Method','Input tokens','Retrievals','Verifications','Seconds'],[[r['dataset'].replace('MultiHopQA',''),r['method'],f(r['metrics']['input_tokens']['mean']),f(r['metrics']['retrieval_calls']['mean']),f(r['metrics']['verification_calls']['mean']),f(r['metrics']['latency']['mean'])] for r in archive],[1.05,1.1,1.25,1.0,1.15,1.15])
    h('6.2 Fresh checkpoint gate and replication stability',2)
    p(f'The checkpoint reward-variance gate {"passed" if gate["passed"] else "failed"}: {gate["n_identical_action_means"]} of {gate["n_checkpoints"]} checkpoints had identical K=5 action-mean EM ({f(gate["fraction"])}). The frozen fallback {"was activated" if counts["checkpoints"]>counts["questions"] else "was not activated"}. This small development gate is not the missing 40-question-per-dataset baseline eligibility study. [E4–E6]')
    rows=[]
    for candidate in ks['candidates']:
        for policy,cell in candidate['by_continuation'].items():rows.append([candidate['K'],policy,f(cell['argmax_agreement']),f(cell['precision_fraction']),'pass' if candidate['passed'] else 'fail'])
    cap(f'Table 4. Independent A/B stability on {counts["checkpoints"]} development checkpoints from {counts["questions"]} questions across the three datasets; Qwen3-4B Q4_K_M, prefix seed 17. Complete seed matrix, top-two intervals, and tie convention are retained in E4–E6. A row’s candidate passes only when both continuations meet both thresholds.')
    table(['K','Continuation','Argmax agreement','Margin precision fraction','Candidate'],rows,[.5,1.2,1.8,2.1,1.1])
    if ks['selected'] is None:p('Neither candidate K satisfied the complete stability rule. The following K=5 results are exploratory measurements, not certified action rankings or frozen settings for a held-out efficacy run. The failure is retained rather than converting nominal confidence intervals into a passed gate. [E5–E6]')
    else:p(f'The rule selected K={ks["selected"]}; both A/B sets are retained. Subsequent held-out inference still requires a trained MIVE controller, complete benchmark manifests, and a separate final freeze. [E4–E6]')
    h('6.3 Continuation discrepancies and rank changes',2)
    metrics=[('signed_THINK_bias','Signed THINK bias'),('absolute_discrepancy','Mean absolute action discrepancy'),('argmax_change','Argmax change'),('kendall_tau_b','Kendall tau b'),('THINK_STOP_rank_reversal','Strict THINK STOP reversal'),('behavioral_disagreement','Behavioral disagreement')]
    cap(f'Table 5. Continuation diagnostics on the fresh development questions from HotpotQA, 2WikiMultiHopQA, and MuSiQue. Qwen3-4B Q4_K_M; prefix seed 17; K={k} per A/B set; {counts["questions"]} question clusters. Brackets are 95% intervals from 10,000 question-bootstrap resamples. Primary EM artifacts: E5–E7; secondary priced-value artifacts: E8. Undefined tau values are excluded with denominators retained in the artifact.')
    table(['Metric','Terminal EM [95% CI]','Priced value [95% CI]'],[[label,interval(summary[key]),interval(secondary[key]) if key in secondary else 'Same policy comparison'] for key,label in metrics],[1.7,2.5,2.5])
    p(f'Tau-b has defined paired ranks in {summary["kendall_tau_b"]["n_questions"]} of {counts["questions"]} question clusters under EM and {secondary["kendall_tau_b"]["n_questions"]} of {counts["questions"]} under priced return. All-tied cases are excluded only from that metric; the other diagnostics retain their reported denominators. [E5, E8]')
    p('The primary EM tau-b interval is degenerate because only one question has defined paired ranks; it does not establish population precision. Argmax changes also depend on the declared tie convention. These features, together with the failed K gate, limit interpretation of apparent rank disagreement.')
    figure('continuation_values',f'Figure 3. Feasible action-mean terminal EM under the labeling and legacy deployment continuations. HotpotQA, 2WikiMultiHopQA, and MuSiQue development split; Qwen3-4B Q4_K_M; {counts["questions"]} questions, {counts["checkpoints"]} checkpoints; one prefix seed (17); two disjoint sets of K={k} continuation replicates. Blank cells indicate infeasible actions. These finite-K values are estimates, with stability reported in Table 4. [E4–E7]')
    figure('mismatch_intervals',f'Figure 4. Continuation discrepancy estimates on terminal EM and secondary resource-priced return. Same three development datasets and Qwen3-4B Q4_K_M; {counts["questions"]} question clusters; prefix seed 17; two independent K={k} sets per continuation. Intervals use 10,000 question-cluster bootstrap resamples and condition on observed branch estimates. Strict EM THINK STOP reversal is structurally impossible under deterministic STOP. [E5–E8]')
    p(f'The strict priced THINK-versus-STOP reversal frequency was {interval(secondary["THINK_STOP_rank_reversal"])}; reversal relative to any alternative action was {interval(secondary["THINK_any_action_rank_reversal"])}. These measured finite-K frequencies connect the constructive example to the retained audit, but do not identify the cause of the legacy held-out allocation failure. Ties are excluded from strict reversal. [E8]')
    h('6.4 Sampling noise and randomness coupling',2)
    cap('Table 6. Noise-versus-shift estimates and equal-branch-budget contrast variances for the fresh audit. Same three development datasets, Qwen3-4B Q4_K_M, one prefix seed, and question-cluster intervals. Squared shift is unbiased under the independent-replicate assumptions and can be negative. E5–E8 retain per-checkpoint values and bootstrap denominators.')
    selected=[('observed_squared_discrepancy','Observed squared discrepancy'),('estimated_noise','Estimated sampling noise'),('estimated_squared_shift','Estimated squared shift'),('independent_gap_variance','Independent THINK contrast variance'),('matched_gap_variance','Matched THINK contrast variance')]
    table(['Metric','Terminal EM [95% CI]','Priced value [95% CI]'],[[label,interval(summary[key]),interval(secondary[key])] for key,label in selected],[2.0,2.35,2.35])
    figure('noise_and_coupling',f'Figure 5. Sampling noise and coupling audit on HotpotQA, 2WikiMultiHopQA, and MuSiQue development checkpoints with Qwen3-4B Q4_K_M. Left: primary-EM squared-discrepancy components with 10,000 question-bootstrap intervals. Right: independent versus matched-seed THINK-contrast variance for the priced objective, with equal K={k} branches per arm. Points are checkpoints from {counts["questions"]} questions; prefix seed 17. The diagonal marks equal variance; shared seeds need not reduce variance. [E5–E8]')
    sensitivity=load(DEV/'runtime_sensitivity.json')
    clean=sensitivity['analyses']['priced_value']
    p(f'In {len(sensitivity["mixed_runtime_checkpoints"])} of {counts["checkpoints"]} checkpoints, the retained records span both serial and four-slot execution. Their outcomes remain in the complete audit, but execution regime can confound a coupling contrast. A separate homogeneous-runtime sensitivity retains {clean["checkpoints"]} checkpoints and estimates matched-minus-independent priced contrast variance as {interval(clean["statistics"]["matched_minus_independent"])}. Its complete primary and secondary statistics and exclusion identities are retained in runtime_sensitivity.json beside E5. This sensitivity changes neither the primary outcomes nor their reported denominator.')
    h('6.5 Verification and offline resource use',2)
    p(f'The fresh reconstruction verified {verification["branches"]} branch outcomes, restored state digests, feasible initial interventions, normalized answer scores, future resource totals, disjoint primary seeds, and the complete replicate matrix. It found {verification["budget_violations"]} budget violations. Across retained server logs, collection used {completion["server_recorded_generation_calls_all_attempts"]} completed physical generation calls and {completion["server_recorded_output_tokens_all_attempts"]} decoded tokens, with {f(completion["server_recorded_generation_seconds_all_attempts"])} summed seconds of server prompt evaluation and decoding. These totals include prefixes, baselines, discarded setup work, and completed requests in the interrupted serial invocation; they exclude startup, transport, non-generation work, and any request without a completed timing record. Concurrent durations are summed, not wall-clock elapsed time. Physical labeling cost must not be equated with deployment-only savings. [E7, E9]')
    p('Recorded server totals are lower bounds on overall generation work if an interrupted request lacks a completion record. The public numeric_server_timings.json ledger retains each completed timing block and source-log hashes, making the reported sums reconstructable without raw benchmark text. [E9]')
    audit_cost=load(DEV/'offline_audit_cost.json');charged=audit_cost['resources']['generated_tokens']
    p(f'Completed branches incurred exactly {int(charged["observed_total"])} future generated-token charges. Mean branch-label cost per question was {f(charged["mean_per_question"])} tokens, with question-bootstrap interval [{f(charged["mean_per_question_ci95"][0])}, {f(charged["mean_per_question_ci95"][1])}]. This interval describes the mean across questions, not uncertainty in the observed total. Prefixes, baselines, and discarded work are excluded from this branch-only total and included where available in the server totals above. All resource coordinates are retained in offline_audit_cost.json beside E5. This is diagnostic labeling cost, not the training cost of a newly evaluated MIVE controller.')
    p('The complete retained reconstruction also checked 642 answer-score and episode-counter records, 296 summary metrics with intervals, 46 paired comparisons, 192 branch payloads, and 1080 model predictions. That audit establishes numerical consistency under the implemented scorer. It does not certify official benchmark equivalence or rerun every historical stochastic generation. [E2]')

    h('7 Discussion')
    p('The principal retained finding is that enforcing a hard resource cap did not make this learned controller efficient. Legacy CRAVE repeatedly reasoned while confidence usually stopped after one candidate. Its additional generation was spread across questions, whereas the observed accuracy differences were concentrated. [E1–E3] This is a failure case for the particular controller and training scale, not a general impossibility result for learned allocation.')
    p('The fresh audit makes continuation sensitivity measurable from restored states. Independent A/B sets expose label instability that a single matched matrix can hide. A failed K gate limits the interpretation of the measured ranks: disagreement can reflect continuation shift, sampling error, or both. The noise decomposition is useful precisely because raw disagreement alone cannot distinguish those explanations. [E4–E8]')
    p('Theorem 3 connects a decision margin to the number of independent replicates required to resolve it. Narrow margins can require far more than a few rollouts. A useful refreshed controller must therefore demonstrate reliable action rankings, not merely complete its training code. Uniform continuation bounds are stronger than average observed discrepancies and should not be claimed from this development sample.')
    p('Offline supervision must be amortized against measured online savings in the same resource unit. If additional labeling and training cost is L and mean deployment savings is S, the positive break-even count is L/S only when S is positive. The retained confidence comparison has negative generated-token savings for CRAVE, so no finite positive token break-even count follows. An interval crossing zero would likewise prevent a bounded upper break-even estimate. This conclusion concerns the observed legacy deployment costs; it supplies no break-even result for an untrained refreshed policy. [E1–E3]')
    p('A future matched-budget comparison with Adaptive-RAG and FLARE must charge their classifier or discarded draft generations and retrieval costs. Self-RAG requires a separate native trained-model track. The comparison matrix in Table 1 prevents simplified proxies from inheriting the names or reported performance of those methods.')
    h('8 Limitations')
    p('The available empirical scope is a small quantized backbone, bounded supplied-context exports, a single retained held-out cap, and a small fresh development audit. Question-cluster intervals are unstable with few independent questions and condition on finite branch estimates. Public labeled validation questions do not establish deployment distribution generalization. Same-backbone verification may repeat an earlier error; reported confidence is not calibrated merely because it appears in JSON. Latency cannot be compared as a cold-runtime benchmark across cached and fresh records.')
    p('The additional backbone checks in Appendix D verify execution only. A powered held-out comparison of a newly trained MIVE controller, faithful published baselines, component ablations, and noninferiority with savings remains future work. Their absence limits causal attribution and deployment recommendations. Appendix B specifies the evidence required to extend the current claims.')
    h('9 Conclusion')
    p('Matched intervention values compare executable actions only after the state, resources, and continuation policy are specified. The retained CRAVE controller obeyed the implemented caps but used more generated tokens than confidence without a reliable accuracy gain. The fresh independent continuation audit reports its gates, finite-K discrepancies, and coupling behavior explicitly. The theoretical results separate execution guarantees from value accuracy and quantify the sampling margin needed for a reliable THINK-versus-STOP decision. A beneficial refreshed controller remains an empirical question requiring a new frozen held-out study.')

    h('Declarations')
    declarations=ART/'author_declarations.json'
    decl=load(declarations) if declarations.exists() else {}
    p('Funding. '+decl.get('funding','Author confirmation required before submission.'))
    p('Competing interests. '+decl.get('competing_interests','Author confirmation required before submission.'))
    p('Ethics. This study uses existing question-answering datasets and does not collect new participant data. No human-subject approval is claimed. Upstream dataset licenses and lawful reconstruction requirements apply; raw benchmark text is excluded from the prepared public archive pending redistribution verification. The author must confirm institutional ethics requirements before submission.')
    p('AI assistance. An AI assistant assisted with implementation, experiment orchestration, numerical reconstruction, reference checks, figure preparation, and manuscript editing. The author remains responsible for the scientific claims, proofs, source attribution, declarations, and final submission. AI assistance does not replace author verification.')
    p('Author contributions. To be confirmed by the author before submission.')
    p('Code and data availability. The legacy numerical release is public at https://github.com/Al1ranger/CRAVE, commit 03a6f228b4c65e7faeb428164bb6771a993da708. The accompanying revision includes audit code, manifests, numerical outcomes, and reconstruction instructions. A permanent archive DOI is pending deposition. Raw questions and passages require lawful access to the stated upstream releases.')

    h('Appendix A Evidence paths')
    p('Artifact IDs in text and captions resolve to the following exact retained paths. Raw values retain full precision; manuscript tables use four decimal places for display only. Recomputed intervals in E3 use the stated new bootstrap seed; original historical tests remain in E1. No raw score has been rounded, replaced, or edited.')
    paths={
      'E1':ROOT/'crave/runs/expanded_8484164aef776a12/evaluation.json',
      'E2':ROOT/'tmp/full_manuscript_assets/comprehensive_result_audit.json',
      'E3':ART/'archival_reanalysis.json','E4':DEV/'freeze.json','E5':DEV/'summary.json',
      'E6':DEV/'stability.json','E7':DEV/'independent_verification.json','E8':DEV/'utility_summary.json',
      'E9':ART/'experiment_register.json'}
    for key,path in paths.items():p(key+'  '+path.as_posix(),'Body Text')
    p('E1 supporting manifests and original statistical tests are in the adjacent FROZEN_EXPANDED_CONFIG.json, task_manifest.json, and results directory. Fresh raw branches, saved checkpoints, baseline trajectories, K-specific estimates, runtime_amendment.json, and the secondary analysis declaration are adjacent to E4. Source-level verification and reference metadata are retained under experiments/mive_revision/artifacts. The experiment register gives executable commands, seeds, hardware, failures, and artifact identities.')
    h('Appendix B Study completion and claim restrictions')
    cap('Table B1. Evidence requirements separated from measured results. This replaces the former main-text status table. Status is evidence availability, not a performance result.')
    table(['Requested study','Current evidence','Narrower honest claim'],[
      ['Checkpoint variance and two continuations','Fresh development audit; see Tables 4–6','Finite-K local sensitivity, conditional on reported gates'],
      ['Scaled 7B and 1.5B evaluation','Official model manifests; execution checks in Appendix D','No held-out efficacy or model-size transfer claim'],
      ['300 questions × 3 datasets × 3 seeds × 3 caps','Prospective intent only; no final freeze','Legacy 40-question, one-cap result only'],
      ['New MIVE held-out controller','No verified refreshed training/evaluation package','MIVE estimand and executable audit'],
      ['Adaptive-RAG and FLARE reproduction','Comparison conditions only','Related-work comparison, not reproduced scores'],
      ['Mask, penalty, value, refresh and tape ablations','Coupling audit only; other requested ablations incomplete','No causal attribution to optional components'],
      ['NI, savings, dev power and break-even intervals','Legacy contrasts only; no MIVE dev power study','No NI claim; legacy token break-even absent'],
      ['Zenodo DOI','Not deposited','Local archive preparation; DOI pending']],[1.8,2.25,2.65])
    p('The intended primary grid requires 32,400 episodes for four policies before the additional published baselines or ablations. This is a prospective count, not an achieved run. The prospective_scaled_plan.json file leaves held-out inference disabled until development choices, official manifests, a trained MIVE controller, and power calculations are complete. It is explicitly not a final experimental freeze. [E9]')
    p('Unsupported claims to delete include the missing 1.5B baseline scores and selected threshold, the unretained six-episode CPU execution, the missing compact continuation-refresh training ledger, measured generalization to 7B models, an efficiency advantage of refreshed MIVE, a multi-cap Pareto frontier, official benchmark evaluator equivalence, and successful noninferiority. The manuscript makes none of these empirical claims.')
    h('Appendix C Archival interpretation')
    p('The original training and post-pilot evaluation remain separate historical evidence. Earlier intermediate-value residuals used only three distinct saved states and cannot substitute for the new independent two-continuation measurement. They are omitted from the primary results rather than presented as an estimate of continuation bias. The retained negative allocation result supports diagnosis; it does not distinguish approximation error, insufficient supervision, continuation mismatch, or inadequate public-state features as causal mechanisms.')
    cap('Table C1. Additional recorded resources and action counts for the retained HotpotQA and 2WikiMultiHopQA evaluation. Qwen3-4B Q4_K_M; 20 questions per dataset; seeds 17, 29, and 43; cap 512 generated tokens. Means come from E3; admitted context and controller time are recorded separately from generated tokens.')
    table(['Dataset','Method','Context tokens','THINK calls','Controller seconds'],[[r['dataset'].replace('MultiHopQA',''),r['method'],f(r['metrics']['context_tokens_added']['mean']),f(r['metrics']['think_calls']['mean']),f(r['metrics']['controller_latency']['mean'])] for r in archive],[1.1,1.2,1.4,1.4,1.6])
    h('Appendix D Additional backbone execution checks')
    smoke_rows=[]
    for model in ['Qwen2.5-1.5B','Qwen2.5-7B']:
        location=ART/('smoke_'+model);result_path=location/'verified_results.json'
        if not result_path.exists():
            p(model+' has no completed checksum-verified execution check in this revision. No numerical result is asserted.')
            continue
        result=load(result_path);assert result['passed']
        frozen=load(location/'freeze.json')
        p(f'{model} Q4_K_M completed {len(result["episodes"])} development episodes using the fixed and confidence policies, one previously frozen question per dataset and seed 17. The confidence threshold was 0.8 and the cap was 512 generated tokens, one retrieval, and one verification. This is an executor and accounting replication, not an efficacy comparison or a continuation audit. Native Qwen2.5 ChatML replaces the Qwen3 thinking directive. Hardware: {frozen["config"]["hardware"]}. The check found {result["budget_violations"]} budget violations. Exact freeze and outcomes: {location.as_posix()}/freeze.json and verified_results.json.')
        for row in result['episodes']:
            smoke_rows.append([model.replace('Qwen2.5-',''),row['dataset'].replace('MultiHopQA',''),row['method'],f(row['exact_match']),f(row['f1']),str(row['output_tokens']),str(row['retrievals'])+'/'+str(row['verifications'])])
    if smoke_rows:
        cap('Table D1. Individual development execution checks on HotpotQA, 2WikiMultiHopQA, and MuSiQue with the listed Qwen2.5 Q4_K_M models. Each dataset contributes one question per model and policy; seed 17 only. These are raw per-episode metrics, with no population confidence intervals or model ranking. R/V denotes charged retrieval and verification calls. Artifacts are identified in Appendix D.')
        table(['Model','Dataset','Policy','EM','F1','Output tokens','R/V'],smoke_rows,[.8,1.25,1.1,.6,.6,1.45,.9])
    h('References')
    reference_audit=load(ART/'references/audit.json');by_number={r['number']:r for r in reference_audit}
    for para in source.paragraphs:
        match=re.match(r'^\[(\d+)\]\s*(.*)',para.text)
        if not match:continue
        number=int(match.group(1));text=para.text.replace('etal.','et al. ').replace('et al.','et al. ')
        text=re.sub(r'\s+',' ',text).strip()
        audit=by_number.get(number,{})
        if number==30:text+=' https://www.routledge.com/Constrained-Markov-Decision-Processes/Altman/p/book/9780849303821'
        if number==27 and 'DOI' not in text:text=text.replace('https://','DOI 10.18653/v1/2024.naacl-long.389. https://',1)
        if number==28 and 'DOI' not in text:text=text.replace('https://','DOI 10.18653/v1/2023.emnlp-main.495. https://',1)
        if number==32 and 'DOI' not in text:text=text.replace('https://','DOI 10.18653/v1/2022.naacl-main.272. https://',1)
        q=p(text);q.paragraph_format.space_after=Pt(5)
        for run in q.runs:run.font.size=Pt(9)
    q=p('[35] Wassily Hoeffding. Probability Inequalities for Sums of Bounded Random Variables. Journal of the American Statistical Association, 58(301):13–30, 1963. DOI 10.1080/01621459.1963.10500830. https://doi.org/10.1080/01621459.1963.10500830')
    for run in q.runs:run.font.size=Pt(9)
    doc.core_properties.title=doc.paragraphs[0].text
    doc.core_properties.author='Ehsan Derakhshi Ghazijahani'
    doc.core_properties.subject='Evidence-led revision with independent development continuation audit'
    doc.save(TARGET)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_sha256
    report={'path':str(TARGET),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':len(doc.inline_shapes),
            'native_math':len(doc.element.findall('.//{http://schemas.openxmlformats.org/officeDocument/2006/math}oMath')),
            'source':str(SOURCE),'source_sha256':source_sha256,'input_preserved':True,'requires_render_verification':True}
    (ART/'manuscript_structure.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report))

if __name__=='__main__':main()
